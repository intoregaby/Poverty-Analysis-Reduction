"""Canonical report AST and its three renderers (EP-14).

Every number shown in HTML, PDF or XLSX is a string that was fixed when the AST was built from
compiled claims and fact tables (`claims.fmt`). Renderers only lay that out; `extract_*` + `parity`
read the produced artefacts back and compare them with the AST, so cross-format drift is detectable.
"""
from __future__ import annotations

import html
import io
import re
import zipfile
import copy

from . import claims as cl
from .contracts import canonical_hash

NUMTOK = re.compile(r"-?\d+(?:\.\d+)?")


def build_ast(report_id: str, version: str, title: str, status: str, sections: list[dict], warnings: list[str],
              verification_headline: str) -> dict:
    """sections: [{id,title,claims:[compiled claim dicts], tables:[{id,title,columns,rows(strings)}], charts:[{id,title,table_id,x,y}]}]"""
    ast = {"report_id": report_id, "version": version, "title": title, "status": status, "sections": copy.deepcopy(sections),
           "warnings": list(warnings), "verification_headline": verification_headline}
    for s in sections:
        for ch in s.get("charts", []):
            if ch["table_id"] not in {t["id"] for t in s.get("tables", [])}:
                raise ValueError(f"chart {ch['id']} references missing table {ch['table_id']}")
    ast["ast_sha256"] = canonical_hash({k: v for k, v in ast.items() if k != "ast_sha256"})
    return ast


def validate_snapshot(ast: dict) -> None:
    expected = canonical_hash({k: v for k, v in ast.items() if k != "ast_sha256"})
    if ast.get("ast_sha256") != expected:
        raise ValueError("report snapshot was modified after freezing")


def ast_numbers(ast: dict) -> list[str]:
    """Every numeric token the report is allowed to display, in document order."""
    out = []
    for s in ast["sections"]:
        for c in s["claims"]:
            out += NUMTOK.findall(c["text"])
        for t in s.get("tables", []):
            for r in t["rows"]:
                for v in r[1:]:
                    out += NUMTOK.findall(str(v))
    return out


def ast_claim_ids(ast: dict) -> list[str]:
    return [c["claim_id"] for s in ast["sections"] for c in s["claims"]]


def _status_banner(ast):
    return {"published": "OFFICIAL", "approved": "APPROVED - NOT YET PUBLISHED"}.get(ast["status"], "DRAFT - NOT OFFICIAL")


# ---- HTML -----------------------------------------------------------------------------------------------------
def render_html(ast: dict) -> bytes:
    validate_snapshot(ast)
    e = html.escape
    p = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>{e(ast['title'])}</title></head><body>",
         f"<header><h1>{e(ast['title'])}</h1><p class='status' data-status='{e(ast['status'])}'>{e(_status_banner(ast))}</p>",
         f"<p class='verification'>{e(ast['verification_headline'])}</p></header>"]
    for w in ast["warnings"]:
        p.append(f"<aside class='warning'>{e(w)}</aside>")
    for s in ast["sections"]:
        p.append(f"<section id='{e(s['id'])}'><h2>{e(s['title'])}</h2>")
        for c in s["claims"]:
            p.append(f"<p class='claim' data-claim-id='{e(c['claim_id'])}'>{e(c['text'])}</p>")
        for t in s.get("tables", []):
            p.append(f"<table id='{e(t['id'])}'><caption>{e(t['title'])}</caption><thead><tr>"
                     + "".join(f"<th>{e(h)}</th>" for h in t["columns"]) + "</tr></thead><tbody>")
            for r in t["rows"]:
                p.append("<tr>" + "".join(f"<td>{e(str(v))}</td>" for v in r) + "</tr>")
            p.append("</tbody></table>")
        p.append("</section>")
    p.append(f"<footer data-ast-sha256='{ast['ast_sha256']}'>Report {e(ast['report_id'])} v{e(ast['version'])}</footer></body></html>")
    return "".join(p).encode("utf-8")


def extract_html(b: bytes) -> dict:
    t = b.decode("utf-8")
    claims = re.findall(r"data-claim-id='([^']+)'>(.*?)</p>", t)
    cells = re.findall(r"<td>(.*?)</td>", t)
    nums = []
    for _, text in claims:
        nums += NUMTOK.findall(html.unescape(text))
    return {"claim_ids": [c for c, _ in claims], "claim_numbers": nums, "table_cell_numbers":
            [n for c in cells for n in NUMTOK.findall(html.unescape(c))],
            "status": re.search(r"data-status='([^']+)'", t).group(1),
            "ast_sha256": re.search(r"data-ast-sha256='([0-9a-f]+)'", t).group(1)}


# ---- PDF ---------------------------------------------------------------------------------------------------------
def render_pdf(ast: dict) -> bytes:
    validate_snapshot(ast)
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table
    st = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=ast["title"], author="Ireme Insight")
    fl = [Paragraph(html.escape(ast["title"]), st["Title"]), Paragraph(html.escape(_status_banner(ast)), st["Heading3"]),
          Paragraph(html.escape(ast["verification_headline"]), st["BodyText"])]
    for w in ast["warnings"]:
        fl.append(Paragraph("WARNING: " + html.escape(w), st["BodyText"]))
    for s in ast["sections"]:
        fl.append(Paragraph(html.escape(s["title"]), st["Heading2"]))
        for c in s["claims"]:
            fl.append(Paragraph(f"[{html.escape(c['claim_id'])}] " + html.escape(c["text"]), st["BodyText"]))
        for t in s.get("tables", []):
            fl.append(Paragraph(html.escape(t["title"]), st["Heading4"]))
            fl.append(Table([t["columns"]] + [[str(v) for v in r] for r in t["rows"]]))
            fl.append(Spacer(1, 8))
    fl.append(Paragraph(f"ast-sha256: {ast['ast_sha256']}", st["Code"]))
    fl.append(Paragraph(f"report-status: {html.escape(ast['status'])}", st["Code"]))
    doc.build(fl)
    return buf.getvalue()


def extract_pdf(b: bytes) -> dict:
    from pypdf import PdfReader
    text = re.sub(r"\s+", " ", " ".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(b)).pages))
    claim_ids = re.findall(r"\[([A-Za-z0-9_.\-]+)\]", text)
    return {"claim_ids": claim_ids, "text": text, "ast_sha256": (re.search(r"ast-sha256:\s*([0-9a-f]{64})", text).group(1)),
            "status": re.search(r"report-status:\s*([a-z_]+)", text).group(1)}


# ---- XLSX with native charts (OPS-05) ---------------------------------------------------------------------------------------
def render_xlsx(ast: dict) -> bytes:
    validate_snapshot(ast)
    import xlsxwriter
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False,
                                   "strings_to_numbers": False})
    meta = wb.add_worksheet("Report")
    meta.write_string(0, 0, ast["title"])
    meta.write_string(1, 0, "status"); meta.write_string(1, 1, ast["status"])
    meta.write_string(2, 0, "verification"); meta.write_string(2, 1, ast["verification_headline"])
    meta.write_string(3, 0, "ast_sha256"); meta.write_string(3, 1, ast["ast_sha256"])
    r = 5
    for s in ast["sections"]:
        for c in s["claims"]:
            meta.write_string(r, 0, c["claim_id"]); meta.write_string(r, 1, _safe(c["text"])); r += 1
    for s in ast["sections"]:
        for t in s.get("tables", []):
            ws = wb.add_worksheet(t["id"][:31])
            for j, h in enumerate(t["columns"]):
                ws.write_string(0, j, h)
            for i, row in enumerate(t["rows"], start=1):
                ws.write_string(i, 0, _safe(str(row[0])))
                for j, v in enumerate(row[1:], start=1):
                    m = NUMTOK.fullmatch(str(v).replace("%", "").strip())
                    ws.write_number(i, j, float(m.group(0))) if m else ws.write_string(i, j, _safe(str(v)))
            for ch in s.get("charts", []):
                if ch["table_id"] != t["id"]:
                    continue
                n = len(t["rows"])
                chart = wb.add_chart({"type": "column"})
                col = t["columns"].index(ch["y"])
                chart.add_series({"name": [ws.get_name(), 0, col], "categories": [ws.get_name(), 1, 0, n, 0],
                                  "values": [ws.get_name(), 1, col, n, col]})
                chart.set_title({"name": ch["title"]})
                ws.insert_chart(1, len(t["columns"]) + 1, chart)
    wb.close()
    return buf.getvalue()


def _safe(s: str) -> str:
    return "'" + s if s.startswith(("=", "+", "-", "@")) and not NUMTOK.fullmatch(s) else s


def extract_xlsx(b: bytes) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True)
    meta = wb["Report"]
    out = {"title": meta["A1"].value, "status": meta["B2"].value, "ast_sha256": meta["B4"].value, "claim_ids": [],
           "claim_numbers": [], "tables": {}}
    for row in meta.iter_rows(min_row=6, values_only=True):
        if row[0]:
            out["claim_ids"].append(row[0])
            out["claim_numbers"] += NUMTOK.findall(str(row[1]))
    for ws in wb.worksheets[1:]:
        out["tables"][ws.title] = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    z = zipfile.ZipFile(io.BytesIO(b))
    out["charts"] = []
    for n in z.namelist():
        if n.startswith("xl/charts/chart"):
            x = z.read(n).decode()
            out["charts"].append({"refs": re.findall(r"<c:f>(.*?)</c:f>", x), "native": "<c:barChart>" in x})
    return out


def parity(ast: dict, html_b: bytes, pdf_b: bytes, xlsx_b: bytes) -> list[str]:
    """Returns discrepancies between the AST and each rendered artefact (empty list = consistent)."""
    validate_snapshot(ast)
    problems = []
    h, p, x = extract_html(html_b), extract_pdf(pdf_b), extract_xlsx(xlsx_b)
    want_ids, want_claim_nums = ast_claim_ids(ast), []
    for s in ast["sections"]:
        for c in s["claims"]:
            want_claim_nums += NUMTOK.findall(c["text"])
    want_table = [n for s in ast["sections"] for t in s.get("tables", []) for r in t["rows"] for v in r[1:] for n in NUMTOK.findall(str(v))]
    if h["claim_ids"] != want_ids or p["claim_ids"] != want_ids or x["claim_ids"] != want_ids:
        problems.append("claim id lists differ between formats")
    if h["claim_numbers"] != want_claim_nums or x["claim_numbers"] != want_claim_nums:
        problems.append("claim numbers differ between formats")
    # Compare the whole PDF's numeric sequence, including repeated values and extra numbers.
    # Membership alone would accept swapped estimates or a fabricated extra percentage.
    visible = [ast["title"], _status_banner(ast), ast["verification_headline"]]
    visible.extend("WARNING: " + w for w in ast["warnings"])
    for section in ast["sections"]:
        visible.append(section["title"])
        visible.extend(f"[{c['claim_id']}] {c['text']}" for c in section["claims"])
        for table in section.get("tables", []):
            visible.extend([table["title"], *table["columns"]])
            visible.extend(str(v) for row in table["rows"] for v in row)
    pdf_body = p["text"].split("ast-sha256:")[0]
    if NUMTOK.findall(pdf_body) != NUMTOK.findall(" ".join(visible)):
        problems.append("PDF numeric sequence differs from AST")
    if h["table_cell_numbers"] != want_table:
        problems.append("HTML table numbers differ from AST")
    xl_clean = [NUMTOK.findall(f"{float(v):.10g}")[0] for ws in x["tables"].values() for r in ws[1:] for v in r[1:] if isinstance(v, (int, float))]
    if [float(n) for n in want_table] != [float(n) for n in xl_clean]:
        problems.append("XLSX table numbers differ from AST")
    for name, got in (("html", h["ast_sha256"]), ("pdf", p["ast_sha256"]), ("xlsx", x["ast_sha256"])):
        if got != ast["ast_sha256"]:
            problems.append(f"{name} AST hash differs")
    if not (h["status"] == x["status"] == p["status"] == ast["status"]):
        problems.append("status differs between formats")
    return problems

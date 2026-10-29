"""Fast checks for catalog coverage, source guards, checkpoints, and reports."""

from __future__ import annotations

import io
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd

from tools.investigation import catalog, engine, finscope2012, sources
from tools.investigation.analysis import weighted_percentage
from tools.investigation.reports import export_reports


class CatalogTests(unittest.TestCase):
    def test_all_supported_columns_and_unsupported_file_are_recorded(self):
        import pyreadstat

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            frame = pd.DataFrame({"district": ["Gisagara", "Nyanza"], "poverty": [1, None], "weight": [2, 1]})
            frame.to_csv(root / "sample.csv", index=False)
            frame.to_json(root / "sample.json", orient="records")
            frame.to_parquet(root / "sample.parquet")
            frame.to_stata(root / "sample.dta", write_index=False)
            pyreadstat.write_sav(frame.fillna(-1), str(root / "sample.sav"))
            frame.to_excel(root / "sample.xlsx", index=False)
            with zipfile.ZipFile(root / "bundle.zip", "w") as archive:
                archive.write(root / "sample.csv", arcname="inside.csv")
            (root / "unreadable.txt").write_text("unsupported", encoding="utf-8")
            result = catalog.discover(root)
            self.assertEqual(result["files_discovered"], 8)
            self.assertEqual(result["skipped_tables"], 1)
            self.assertEqual(result["failed_tables"], 0)
            self.assertEqual(result["variables_inventoried"], 21)
            for table in result["tables"]:
                if table["status"] == "Completed":
                    self.assertEqual({item["name"] for item in table["variables"]}, {"district", "poverty", "weight"})

    def test_historical_poverty_and_identifiers_cannot_be_used_as_direct_estimates(self):
        self.assertEqual(catalog.classify("pov_jan", "Poverty", "table", "EICV5", 0, 10)[0], "Not comparable")
        self.assertEqual(catalog.classify("hhid", "Household ID", "cs_poverty", "EICV7", 0, 10)[0], "Restricted")
        self.assertEqual(catalog.survey_identity("EICV3 and EICV4 Panel.dta")[0], "EICV3–EICV4 panel")

    def test_new_file_changes_data_signature(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "base.csv").write_text("a\n1\n", encoding="utf-8")
            before, _ = engine.data_signature(root)
            previous = catalog.discover(root)
            (root / "new.csv").write_text("district,poverty\nGisagara,1\n", encoding="utf-8")
            after, _ = engine.data_signature(root)
            self.assertNotEqual(before, after)
            with patch.object(catalog, "_tabular_metadata", wraps=catalog._tabular_metadata) as reader:
                updated = catalog.discover(root, previous=previous)
                self.assertEqual(updated["variables_inventoried"], 3)
                self.assertEqual(reader.call_count, 1)


class SourceTests(unittest.TestCase):
    def test_confidence_interval_parser_skips_the_contents_page(self):
        from reportlab.pdfgen.canvas import Canvas
        with tempfile.TemporaryDirectory() as directory:
            file = Path(directory) / "report.pdf"
            pdf = Canvas(str(file))
            pdf.drawString(20, 700, "Table A9.23: Headcount Poverty rate (contents entry)")
            pdf.showPage()
            pdf.drawString(20, 700, "Table A9.23: Headcount Poverty rate")
            pdf.drawString(20, 680, "95%CI lower bound")
            pdf.drawString(20, 660, "Gisagara 45.59 2.69 40.32 50.86 5.89 1.29")
            pdf.save()
            register = {"sources": [{"id": "eicv7-main", "status": "Completed", "local_cache": str(file), "url": "https://statistics.gov.rw"}]}
            interval = sources.published_poverty_intervals(register)["Gisagara"]
            self.assertEqual((interval["lower"], interval["upper"]), (40.32, 50.86))

    def test_failed_update_retains_separately_verified_cache(self):
        from urllib.error import URLError
        import hashlib
        with tempfile.TemporaryDirectory() as directory:
            store = Path(directory)
            cache = store / "source-cache"
            cache.mkdir()
            pdf = cache / "source.pdf"
            pdf.write_bytes(b"%PDF previously verified fixture")
            old = {**sources.SOURCES[1], "status": "Completed", "local_cache": str(pdf), "version": hashlib.sha256(pdf.read_bytes()).hexdigest(), "retrieved_at": "2026-10-24"}
            run = store / "runs" / "previous"
            run.mkdir(parents=True)
            (run / "02-checking-sources.json").write_text(json.dumps({"sources": [old]}))
            with patch.object(sources, "urlopen", side_effect=URLError("offline")):
                result = sources.verify_source(sources.SOURCES[1], cache)
            self.assertEqual(result["status"], "Failed")
            self.assertEqual(result["cached_evidence"]["retrieved_at"], "2026-10-24")
            pdf.write_bytes(b"tampered")
            with patch.object(sources, "urlopen", side_effect=URLError("offline")):
                self.assertNotIn("cached_evidence", sources.verify_source(sources.SOURCES[1], cache))

    def test_non_https_source_is_rejected_without_retrieval(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(sources, "urlopen") as opener:
            result = sources.verify_source({"id": "bad", "url": "http://example.com/report.pdf", "title": "Insecure source"}, Path(directory))
        self.assertEqual(result["status"], "Failed")
        self.assertIn("HTTPS", result["error"])
        opener.assert_not_called()

    def test_redirect_to_another_site_is_rejected(self):
        from tools.investigation.registry import same_site

        self.assertTrue(same_site("https://rdb.rw/", "https://www.rdb.rw/x"))
        self.assertFalse(same_site("https://rdb.rw/", "https://evil.example/x"))

    def test_network_failure_is_recorded(self):
        from urllib.error import URLError

        with tempfile.TemporaryDirectory() as directory, patch.object(sources, "urlopen", side_effect=URLError("offline")):
            result = sources.verify_source(sources.SOURCES[0], Path(directory))
        self.assertEqual(result["status"], "Failed")
        self.assertIn("offline", result["error"])

    def test_conflicting_narrative_is_flagged(self):
        from reportlab.pdfgen.canvas import Canvas

        with tempfile.TemporaryDirectory() as directory:
            file = Path(directory) / "report.pdf"
            pdf = Canvas(str(file))
            pdf.drawString(20, 700, "Gisagara (46.6%)")
            pdf.save()
            register = {"sources": [{"id": "eicv7-main", "status": "Completed", "local_cache": str(file)}]}
            conflicts = sources.detect_published_rate_conflicts(register, [{"district": "Gisagara", "poverty": 45.6}])
            self.assertEqual(len(conflicts), 1)
            self.assertEqual(conflicts[0]["published_narrative"], 46.6)


class AnalysisTests(unittest.TestCase):
    def test_brief_preserves_overall_coverage_and_reports_measured_problem(self):
        from tools.investigation.brief import analyst_brief
        analysis = {"selected": {"district": "Example", "poverty": 40, "grid_connection": 20, "vup_among_poor": 10, "sample_households": 486},
                    "national": {"poverty": 20, "grid_connection": 50}}
        brief = analyst_brief(analysis, {"source_conflicts": []})
        self.assertIn("40.0%", brief["observations"][0]["text"])
        self.assertIn("20.0 percentage points above", brief["observations"][0]["text"])
        self.assertIn("all 486", brief["coverage"])
        self.assertIn("does not remove", brief["coverage"])
        self.assertIn("not enough to authorize", brief["decision"])

    def test_cluster_variance_and_singleton_stratum_guard(self):
        from tools.investigation.context import design_variance, ratio
        frame = pd.DataFrame({"strata_id": [1, 1, 2, 2], "clust": [1, 2, 3, 4], "pop_wt": [1, 1, 1, 1], "pov_jan": ["No", "Poor", "No", "Poor"]})
        estimate, influence = ratio(frame, pd.Series(True, index=frame.index))
        self.assertEqual(estimate, 0.5)
        self.assertAlmostEqual(design_variance(frame, influence), 0.125)
        with self.assertRaises(ValueError):
            design_variance(frame.iloc[:1], influence.iloc[:1])

    def test_weights_and_missing_values(self):
        frame = pd.DataFrame({"weight": [3, 1, 5], "answer": ["Yes", "No", None]})
        self.assertEqual(weighted_percentage(frame, frame.answer == "Yes", "weight", frame.answer.notna()), 75.0)

    def test_invalid_weights_are_rejected(self):
        frame = pd.DataFrame({"weight": [1, -1], "answer": [True, False]})
        with self.assertRaises(ValueError):
            weighted_percentage(frame, frame.answer, "weight")


class AIReviewTests(unittest.TestCase):
    def test_provider_receives_identified_client_header(self):
        from tools.investigation.ai import review_evidence
        with patch.dict("os.environ", {"GROQ_API_KEY": "test-key"}), patch("tools.investigation.ai.urlopen") as connect:
            connect.return_value.__enter__.return_value.read.return_value = json.dumps({"choices": [{"message": {"content": json.dumps({"finding_ids": ["poverty"], "option_titles": [], "question_ids": ["sampling"]})}}]}).encode()
            result = review_evidence({"findings": [{"id": "poverty"}], "source_conflicts": []}, {"options": []})
            self.assertEqual(result["status"], "Completed")
            self.assertEqual(connect.call_args.args[0].get_header("User-agent"), "NISR-Track2-Investigation/1.0")
    def test_financial_trend_reaches_the_model_as_valid_json(self):
        from tools.investigation.ai import review_evidence
        trend = {"status": "Completed", "labels": {"a": "Aggregate comparison only"}, "national": {"y2020": {"excluded": {"pct": 7.4}}, "y2020_restated_published": {"excluded": 7},
                 "y2024_published": {"excluded": 4}, "mobile_money_2024_estimated": 76.2}, "change_counts": {"excluded": {"Clear fall": 5}}, "checks": [], "limitations": ["x"],
                 "changes": [{"district": "Gisagara", "excluded": {"y2020": {"pct": 9.0}}, "informal_only": {}, "formal": {}, "mobile_money": {}}]}
        with patch.dict("os.environ", {"GROQ_API_KEY": "test-key"}), patch("tools.investigation.ai.urlopen") as connect:
            connect.return_value.__enter__.return_value.read.return_value = json.dumps({"choices": [{"message": {"content": json.dumps({"finding_ids": ["poverty"], "option_titles": [], "question_ids": []})}}]}).encode()
            review_evidence({"findings": [{"id": "poverty"}], "source_conflicts": []}, {"options": []}, financial_access={"district": "Gisagara"}, financial_trend=trend)
            body = json.loads(connect.call_args.args[0].data)
            sent = json.loads(body["messages"][1]["content"])
            self.assertEqual(sent["financial_trend"]["districts"][0]["district"], "Gisagara")
            self.assertEqual(sent["financial_trend"]["national"]["y2024_published"]["excluded"], 4)
        with patch.dict("os.environ", {"GROQ_API_KEY": "test-key"}), patch("tools.investigation.ai.urlopen") as connect:
            connect.return_value.__enter__.return_value.read.return_value = json.dumps({"choices": [{"message": {"content": json.dumps({"finding_ids": ["poverty"], "option_titles": [], "question_ids": []})}}]}).encode()
            review_evidence({"findings": [{"id": "poverty"}], "source_conflicts": []}, {"options": []}, financial_trend={"status": "Unavailable", "reason": "x"})
            self.assertIsNone(json.loads(json.loads(connect.call_args.args[0].data)["messages"][1]["content"])["financial_trend"])

    def test_financial_series_reaches_the_model_as_valid_json(self):
        from tools.investigation.ai import review_evidence
        change = {"points": 1.0, "label": "Clear rise", "ci_lower": 0.1, "ci_upper": 2.0}
        series = {"status": "Completed", "labels": {"national": "Aggregate comparison only"}, "series": [{"key": "2016", "excluded": 11.9}], "checks": [], "limitations": ["x"],
                  "provinces": [{"province": "Eastern", "excluded": {"y2016_pct": 10.3, "y2020_pct": 9.7, "change": change}, "formal": {"y2016_pct": 65.0, "y2020_pct": 75.7, "change": None}}]}
        with patch.dict("os.environ", {"GROQ_API_KEY": "test-key"}), patch("tools.investigation.ai.urlopen") as connect:
            connect.return_value.__enter__.return_value.read.return_value = json.dumps({"choices": [{"message": {"content": json.dumps({"finding_ids": ["poverty"], "option_titles": [], "question_ids": []})}}]}).encode()
            review_evidence({"findings": [{"id": "poverty"}], "source_conflicts": []}, {"options": []}, financial_series=series)
            sent = json.loads(json.loads(connect.call_args.args[0].data)["messages"][1]["content"])
            self.assertEqual(sent["financial_series"]["series"][0]["excluded"], 11.9)
            self.assertNotIn("provinces", sent["financial_series"])

    def test_missing_provider_is_explicit(self):
        from tools.investigation.ai import review_evidence
        with patch.dict("os.environ", {"GROQ_API_KEY": "", "NVIDIA_API_KEY": ""}):
            self.assertEqual(review_evidence({}, {})["status"], "Skipped")

    def test_keys_are_matched_to_providers_by_prefix_and_deduplicated(self):
        from tools.investigation.ai import providers_from_env
        with patch.dict("os.environ", {"GROQ_API_KEY": "nvapi-abc", "NVIDIA_API_KEY": "nvapi-abc"}):
            found = providers_from_env()
            self.assertEqual([item["name"] for item in found], ["NVIDIA"])
        with patch.dict("os.environ", {"GROQ_API_KEY": "gsk_x", "NVIDIA_API_KEY": "nvapi-y"}):
            self.assertEqual([item["name"] for item in providers_from_env()], ["Groq", "NVIDIA"])

    def test_env_file_fills_only_missing_keys(self):
        from tools.investigation.ai import load_env_file
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / ".env.local"
            path.write_text("\n".join(['GROQ_API_KEY="from-file"', "NVIDIA_API_KEY=nvapi-file", "OTHER=1"]), encoding="utf-8")
            with patch.dict("os.environ", {"GROQ_API_KEY": "already", "NVIDIA_API_KEY": ""}, clear=False):
                load_env_file(path)
                self.assertEqual(os.environ["GROQ_API_KEY"], "already")
                self.assertEqual(os.environ["NVIDIA_API_KEY"], "nvapi-file")
                self.assertNotIn("OTHER", os.environ)

    def test_rate_limit_is_retried_once_and_daily_limit_is_not(self):
        from urllib.error import HTTPError
        from tools.investigation.ai import call_provider
        provider = {"name": "Groq", "url": "u", "model": "m", "reasoning": True, "key": "k"}
        short = HTTPError("u", 429, "limit", {"retry-after": "2"}, io.BytesIO(b"slow down"))
        with patch("tools.investigation.ai._send", side_effect=[short, {"ok": 1}]) as send:
            waits = []
            self.assertEqual(call_provider(provider, {}, sleep=waits.append), {"ok": 1})
            self.assertEqual((send.call_count, waits), (2, [2.0]))
        daily = HTTPError("u", 429, "limit", {"retry-after": "2"}, io.BytesIO(b"tokens per day (TPD) reached"))
        with patch("tools.investigation.ai._send", side_effect=[daily]) as send, self.assertRaises(HTTPError):
            call_provider(provider, {}, sleep=lambda _: None)
        self.assertEqual(send.call_count, 1)

    def test_second_provider_answers_when_first_fails(self):
        from urllib.error import HTTPError
        from tools.investigation.ai import review_evidence
        good = {"choices": [{"message": {"content": json.dumps({"finding_ids": ["poverty"], "option_titles": [], "question_ids": []})}}]}
        limit = HTTPError("u", 401, "no", {}, io.BytesIO(b""))
        with patch.dict("os.environ", {"GROQ_API_KEY": "gsk_x", "NVIDIA_API_KEY": "nvapi-y"}), patch("tools.investigation.ai._send", side_effect=[limit, good]):
            result = review_evidence({"findings": [{"id": "poverty"}], "source_conflicts": []}, {"options": []})
        self.assertEqual((result["status"], result["provider"]), ("Completed", "NVIDIA"))
        self.assertIn("Groq", result["fallback_after"][0])
        with patch.dict("os.environ", {"GROQ_API_KEY": "gsk_x", "NVIDIA_API_KEY": ""}), patch("tools.investigation.ai._send", side_effect=[limit]):
            result = review_evidence({"findings": [{"id": "poverty"}], "source_conflicts": []}, {"options": []})
        self.assertEqual((result["status"], result["http_status"]), ("Warning", 401))

    def test_numbered_narrative_keys_are_gathered_and_still_validated(self):
        from tools.investigation.ai import validate_review
        findings = {"findings": [{"id": "poverty", "value": 45.6}]}
        sentence = {"text": "Poverty is 45.6% here.", "finding_ids": ["poverty"]}
        value = {"finding_ids": ["poverty"], "option_titles": [], "question_ids": [], "1": sentence, "2": {"text": "Poverty is 99.9% here.", "finding_ids": ["poverty"]}}
        result = validate_review(value, findings, {"options": []}, {"findings": findings["findings"]})
        self.assertEqual([item["text"] for item in result["narrative"]], [sentence["text"]])
        self.assertEqual(len(result["dropped_statements"]), 1)

    def test_invented_evidence_references_are_rejected(self):
        from tools.investigation.ai import validate_review
        with self.assertRaises(ValueError):
            validate_review({"finding_ids": ["invented"], "option_titles": [], "question_ids": []}, {"findings": [{"id": "poverty"}]}, {"options": []})

    def test_valid_review_keeps_all_findings_and_only_known_questions(self):
        from tools.investigation.ai import validate_review
        result = validate_review({"finding_ids": ["poverty"], "option_titles": [], "question_ids": ["sampling"]}, {"findings": [{"id": "poverty"}]}, {"options": []})
        self.assertEqual(result["status"], "Completed")
        self.assertEqual(len(result["questions"]), 1)


class WorkflowTests(unittest.TestCase):
    def test_checkpoint_retry_does_not_duplicate_completed_work(self):
        import uuid

        with tempfile.TemporaryDirectory() as directory, patch.object(engine, "STORE", Path(directory)):
            work = engine.Investigation(str(uuid.uuid4()), "Gisagara")
            calls = []
            def successful():
                calls.append("success")
                return {"count": 1}
            self.assertEqual(work.stage(engine.STAGES[0], "test", successful), {"count": 1})
            self.assertEqual(work.stage(engine.STAGES[0], "test", successful), {"count": 1})
            self.assertEqual(calls, ["success"])
            with self.assertRaises(ValueError):
                work.stage(engine.STAGES[1], "test", lambda: (_ for _ in ()).throw(ValueError("offline")))
            self.assertEqual(work.manifest["stages"][1]["status"], "Failed")
            self.assertEqual(work.stage(engine.STAGES[1], "retry", lambda: {"sources": 0}), {"sources": 0})
            self.assertEqual(work.manifest["stages"][1]["status"], "Completed")
            self.assertEqual(json.loads(work.manifest_path.read_text())["stages"][1]["status"], "Completed")

    def test_aggregate_reports_open_and_excel_chart_references_data(self):
        from openpyxl import load_workbook
        from pypdf import PdfReader

        row = {"district": "Gisagara", "poverty": 45.6, "vup_among_poor": 18.2, "grid_connection": 39.1,
               "health_insurance": 85.0, "sample_households": 486, "sample_poor_households": 194}
        national = {**row, "district": "Rwanda", "poverty": 27.4}
        analysis = {"district": "Gisagara", "survey": "EICV7 2023/24", "selected": row,
                    "comparators": [], "national": national, "method": {"comparators": "Same survey", "uncertainty": "No CI"}, "not_analyzed": []}
        prepared = {"findings": {"findings": [{"title": "Poverty", "value": 45.6, "national": 27.4,
                    "unit": "% of people", "period": "2023/24", "source": "https://statistics.gov.rw", "limitations": "No CI",
                    "explanation": "Poverty is 45.6%", "sample": 486}]}, "plan": {"approval_status": "Draft", "reason": "Review"},
                    "evaluation": {"status": "Awaiting", "note": "No causal claim"}, "research": {"conclusion": "No intervention", "options": []}, "source_conflicts": []}
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            export_reports(folder, {"id": "test", "started_at": "2026-10-25", "status": "Completed", "data_signature": "abc"}, analysis,
                           {"sources": []}, prepared)
            workbook = load_workbook(folder / "investigation.xlsx")
            self.assertEqual(len(workbook.sheetnames), 11)
            self.assertEqual(workbook["Data Tables"]["B2"].value, 45.6)
            self.assertEqual(len(workbook["Editable Charts"]._charts), 1)
            chart = workbook["Editable Charts"]._charts[0]
            self.assertIn("Data Tables", str(chart.series[0].val.numRef.f))
            self.assertGreaterEqual(len(PdfReader(folder / "investigation.pdf").pages), 1)


class DriverTests(unittest.TestCase):
    def synthetic(self):
        import numpy as np
        import pandas as pd
        from tools.investigation.drivers import VARIABLES

        rng = np.random.default_rng(1)
        n, clusters = 3000, 150
        frame = pd.DataFrame({"clust": np.repeat(np.arange(clusters), n // clusters), "strata": np.repeat(np.arange(clusters) % 5, n // clusters)})
        frame["district"] = np.where(frame.clust < 30, "A", "B")
        frame["province"] = "P"
        frame["weight"] = 1.0
        for name, _, _, kind, _ in VARIABLES:
            frame[name] = rng.integers(0, 2, n).astype(float) if kind == "binary" else rng.random(n)
        frame.loc[frame.district == "A", "no_grid"] = (rng.random((frame.district == "A").sum()) < 0.9).astype(float)
        frame["poor"] = (rng.random(n) < 0.15 + 0.3 * frame.no_grid).astype(float)
        frame["fgt1"] = frame.poor * 0.3
        frame["fgt2"] = frame.poor * 0.09
        return frame

    def test_decomposition_is_exact_and_finds_planted_driver(self):
        from tools.investigation.drivers import decompose

        result = decompose(self.synthetic(), "A", replicates=40)
        self.assertAlmostEqual(result["explained_points"] + result["unexplained_points"], result["gap_points"], places=1)
        top = max(result["groups"], key=lambda item: item["contribution_points"])
        self.assertEqual(top["id"], "energy")
        self.assertEqual(top["grade"], "Strong association")
        self.assertIn("A", [row["district"] for row in result["place_effect"]["districts"]])

    def test_small_cells_are_suppressed(self):
        from tools.investigation.drivers import decompose

        frame = self.synthetic()
        frame.loc[frame.district == "A", "head_female"] = 0.0
        frame.loc[frame.index[:5], "head_female"] = 1.0
        row = next(item for item in decompose(frame, "A", replicates=20)["profile"] if item["name"] == "head_female")
        self.assertTrue(row["status"].startswith("Suppressed"))


class ConsumptionViewTests(unittest.TestCase):
    def frame(self):
        frame = DriverTests("test_decomposition_is_exact_and_finds_planted_driver").synthetic()
        rng = np.random.default_rng(5)
        frame["log_consumption"] = 12.0 - 0.8 * frame.no_grid + 0.1 * rng.standard_normal(len(frame))
        return frame

    def test_shortfall_splits_exactly_and_the_planted_driver_is_found(self):
        from tools.investigation.drivers import consumption_view
        result = consumption_view(self.frame(), "A", replicates=40)
        self.assertGreater(result["shortfall_points"], 0, "district A has more households without the grid, so its spending is lower")
        self.assertAlmostEqual(result["explained_points"] + result["unexplained_points"], result["shortfall_points"], places=1)
        top = max(result["groups"], key=lambda item: item["contribution_points"])
        self.assertEqual(top["id"], "energy")
        self.assertEqual(top["grade"], "Strong association")
        self.assertLess(result["spending_difference_pct"], 0)
        self.assertEqual(result["ranked_drivers"][0], "energy")

    def test_a_district_better_off_than_rwanda_gets_a_negative_shortfall(self):
        from tools.investigation.drivers import consumption_view
        frame = self.frame()
        frame.loc[frame.district == "A", "no_grid"] = 0.0
        frame["log_consumption"] = 12.0 - 0.8 * frame.no_grid
        result = consumption_view(frame, "A", replicates=20)
        self.assertLess(result["shortfall_points"], 0)
        self.assertGreater(result["spending_difference_pct"], 0)

    def test_out_of_sample_fit_is_honest_and_higher_for_spending_than_for_poverty_status(self):
        from tools.investigation.drivers import consumption_view
        result = consumption_view(self.frame(), "A", replicates=20)
        fit = result["explainability"]
        self.assertGreater(fit["living_standards_out_of_sample_r_squared"], 0.5)
        self.assertLessEqual(fit["living_standards_out_of_sample_r_squared"], result["weighted_r_squared"] + 0.02)
        self.assertGreater(fit["living_standards_out_of_sample_r_squared"], fit["poverty_status_out_of_sample_r_squared"])
        self.assertEqual(fit["folds"], 5)

    def test_missing_column_and_tiny_districts_are_rejected(self):
        from tools.investigation.drivers import consumption_view
        frame = self.frame()
        with self.assertRaises(ValueError):
            consumption_view(frame.drop(columns=["log_consumption"]), "A", replicates=5)
        frame.loc[frame.district == "A", "district"] = "C"
        frame.loc[frame.index[:50], "district"] = "A"
        with self.assertRaises(ValueError):
            consumption_view(frame, "A", replicates=5)

    def test_out_of_sample_folds_hold_out_whole_clusters(self):
        from tools.investigation import drivers
        frame = self.frame()
        names = [item[0] for item in drivers.VARIABLES]
        design = np.column_stack([np.ones(len(frame)), frame[names].to_numpy(float)])
        first = drivers.out_of_sample_r_squared(frame, design, frame.log_consumption.to_numpy())
        second = drivers.out_of_sample_r_squared(frame, design, frame.log_consumption.to_numpy())
        self.assertEqual(first, second, "the folds are fixed by a seed, so the result must repeat")


class ReliabilityTests(unittest.TestCase):
    def frame(self):
        rows = [(stratum, cluster, "D" if cluster % 2 else "E", 1.0, 1 if (cluster + member) % 3 == 0 else 0)
                for stratum in (1, 2) for cluster in range(stratum * 10, stratum * 10 + 10) for member in range(5)]
        return pd.DataFrame(rows, columns=["strata_id", "clust", "district", "weight", "hit"])

    def test_domain_ratio_matches_weighted_share_and_has_interval(self):
        from tools.investigation.reliability import domain_ratio
        frame = self.frame()
        result = domain_ratio(frame, frame.district == "D", frame.hit == 1, pd.Series(True, index=frame.index), "weight")
        expected = frame[frame.district == "D"].hit.mean()
        self.assertAlmostEqual(result["estimate"], expected)
        self.assertGreater(result["se"], 0)
        self.assertEqual(result["n"], int((frame.district == "D").sum()))

    def test_invalid_weights_are_rejected(self):
        from tools.investigation.reliability import domain_ratio
        frame = self.frame()
        frame.loc[0, "weight"] = 0
        with self.assertRaises(ValueError):
            domain_ratio(frame, frame.district == "D", frame.hit == 1, pd.Series(True, index=frame.index), "weight")

    def test_grades_follow_documented_rule(self):
        from tools.investigation.reliability import grade
        self.assertEqual(grade(0.5, 0.01, 29), "Suppressed")
        self.assertEqual(grade(0.5, 0.05, 100), "Reliable")
        self.assertEqual(grade(0.5, 0.10, 100), "Use with caution")
        self.assertEqual(grade(0.1, 0.05, 100), "Low reliability")
        self.assertEqual(grade(None, None, 100), "Suppressed")

    def test_national_reconciliation_reports_every_mismatch(self):
        from tools.investigation.reliability import reconcile_national
        result = reconcile_national({"poverty": 27.4, "health_insurance": 80.0}, {"poverty": 27.4, "health_insurance": 85.3})
        self.assertEqual([item["matches"] for item in result], [True, False])


class FinScopeTests(unittest.TestCase):
    def frame(self):
        from tools.investigation.finance import CODES, CORE_COLUMNS as COLUMNS
        rng_rows = []
        for index in range(240):
            district = ["Alpha", "Beta", "Gamma", "Delta", "Eps", "Zeta", "Eta", "Theta", "Iota", "Kappa"][index % 10]
            row = {"a1": "East", "a2": district, "a6": "Rural", "pop_wt": 10.0 + index % 7, "b1": 30}
            bank = index % 4 == 0 and index % 10 < 5
            informal = index % 3 == 0
            for code in CODES:
                used = (code == "01" and bank) or (code == "05" and index % 2 == 0) or (code == "10" and informal)
                row[f"qf1_{code}"] = "Yes" if used else "No"
                row[f"qf4_{code}"] = "Yes" if used else None
            row["qf1_12"] = "No"
            rng_rows.append(row)
        return pd.DataFrame(rng_rows)[COLUMNS]

    def test_strand_precedence_and_mutual_exclusivity(self):
        from tools.investigation.finance import assign_strand
        frame = self.frame()
        strand = assign_strand(frame)
        self.assertTrue(strand.isin(["banked", "other_formal", "informal_only", "excluded"]).all())
        both = frame.index[(frame.qf4_01 == "Yes") & (frame.qf4_05 == "Yes")]
        self.assertTrue(len(both) > 0 and (strand[both] == "banked").all())

    def test_strand_shares_sum_to_one_hundred(self):
        from tools.investigation.finance import finscope_estimates
        result = finscope_estimates(self.frame(), "Alpha", {name: 20.0 + index for index, name in enumerate(["Alpha", "Beta", "Gamma", "Delta", "Eps", "Zeta", "Eta", "Theta", "Iota", "Kappa"])})
        total = sum(item["pct"] for item in result["national"]["strands"].values())
        self.assertAlmostEqual(total, 100.0, delta=0.3)
        self.assertEqual(result["association_with_poverty"]["districts_compared"], 10)
        self.assertEqual(result["reconciled_with_published"], "Approximate")
        self.assertEqual(set(result["reconciliation_with_published"]["national"]), {"banked", "other_formal", "informal_only", "excluded"})

    def test_someone_elses_account_and_mutuelle_rules(self):
        from tools.investigation.finance import assign_strand, MOBILE_MONEY_USERS
        frame = self.frame()
        for column in ["k1", "k3", "i5a_3"]:
            frame[column] = "No"
        frame["f5"] = "You have not heard about the Mobile Money"
        blank = frame.index[(frame.qf4_01.isna()) & (frame.qf4_05.isna()) & (frame.qf4_10.isna())]
        self.assertTrue(len(blank) > 0)
        target = blank[0]
        self.assertEqual(assign_strand(frame)[target], "excluded")
        frame.loc[target, "k3"] = "Yes"
        self.assertEqual(assign_strand(frame)[target], "banked")
        frame.loc[target, "k3"] = "No"
        frame.loc[target, "i5a_3"] = "Yes"
        self.assertEqual(assign_strand(frame)[target], "excluded", "Mutuelle de Santé must not make someone formally served")
        frame.loc[target, "f5"] = MOBILE_MONEY_USERS[3]
        self.assertEqual(assign_strand(frame)[target], "other_formal")

    def test_health_insurance_never_changes_the_strand(self):
        from tools.investigation.finance import assign_strand
        frame = self.frame()
        before = assign_strand(frame)
        frame["i5a_3"] = ["Yes" if index % 5 == 0 else "No" for index in range(len(frame))]
        frame["i5b_1"] = "Yes"
        self.assertTrue((assign_strand(frame) == before).all())

    def test_published_district_reference_is_complete(self):
        from tools.investigation.finance import PUBLISHED_DISTRICTS
        self.assertEqual(len(PUBLISHED_DISTRICTS), 30)
        for name, values in PUBLISHED_DISTRICTS.items():
            known = [value for value in values if value is not None]
            self.assertTrue(98 <= sum(known) <= 102, name)

    def test_reconcile_reports_the_largest_district_gap_and_never_hides_it(self):
        from tools.investigation.finance import PUBLISHED_DISTRICTS, reconcile
        national = {"strands": {key: {"pct": value} for key, value in {"banked": 22.0, "other_formal": 65.0, "informal_only": 5.0, "excluded": 8.0}.items()}}
        rows = []
        for name, (banked, other, informal, excluded) in PUBLISHED_DISTRICTS.items():
            rows.append({"district": name, "banked_pct": banked, "other_formal_pct": other, "informal_only_pct": informal,
                         "excluded_pct": (excluded if excluded is not None else 0) + (17.0 if name == "Nyaruguru" else 1.0)})
        result = reconcile(national, rows)
        self.assertEqual(result["districts"]["largest_difference"]["excluded"], {"points": 17.0, "district": "Nyaruguru"})
        self.assertIn("up to 17.0 points (Nyaruguru)", result["reading"])
        self.assertIn("official", result["reading"])

    def test_selected_district_carries_nisr_published_strand(self):
        from tools.investigation.finance import PUBLISHED_DISTRICTS
        for name, values in PUBLISHED_DISTRICTS.items():
            self.assertEqual(len(values), 4, name)

    def test_unknown_district_and_bad_weights_are_rejected(self):
        from tools.investigation.finance import finscope_estimates
        with self.assertRaises(ValueError):
            finscope_estimates(self.frame(), "Nowhere", {})
        frame = self.frame()
        frame.loc[0, "pop_wt"] = 0
        with self.assertRaises(ValueError):
            finscope_estimates(frame, "Alpha", {})

    def test_current_use_without_ever_use_is_rejected(self):
        from tools.investigation.finance import finscope_estimates
        frame = self.frame()
        frame.loc[1, "qf4_02"] = "Yes"
        frame.loc[1, "qf1_02"] = "No"
        with self.assertRaises(ValueError):
            finscope_estimates(frame, "Alpha", {})

    def test_unmatched_districts_are_reported_and_association_needs_enough_districts(self):
        from tools.investigation.finance import finscope_estimates
        result = finscope_estimates(self.frame(), "Alpha", {"Alpha": 30.0, "Beta": 40.0})
        self.assertEqual(result["association_with_poverty"]["status"], "Not estimated")
        self.assertEqual(len(result["unmatched_districts"]), 8)


class TrendTests(unittest.TestCase):
    def frame(self, shift=0.0):
        rows = [(cluster, "A" if cluster < 20 else "B", 1.0 + (index % 3), index % 5 < (2 if cluster < 20 else 3) + shift)
                for cluster in range(40) for index in range(10)]
        return pd.DataFrame(rows, columns=["clust", "district", "pop_wt", "poor"])

    def test_estimate_matches_weighted_share_and_rejects_bad_weights(self):
        from tools.investigation.trends import estimate
        frame = self.frame()
        domain = frame.district == "A"
        result = estimate(frame, domain, "poor")
        group = frame[domain]
        self.assertAlmostEqual(result["p"], float((group.pop_wt * group.poor).sum() / group.pop_wt.sum()))
        self.assertGreater(result["se"], 0)
        frame.loc[0, "pop_wt"] = 0
        with self.assertRaises(ValueError):
            estimate(frame, domain, "poor")

    def test_small_domains_are_withheld(self):
        from tools.investigation.trends import estimate, shown
        frame = self.frame()
        tiny = pd.Series(False, index=frame.index)
        tiny.iloc[:5] = True
        self.assertIsNone(shown(estimate(frame, tiny, "poor")).get("pct"))

    def test_holm_is_monotone_and_never_below_raw(self):
        from tools.investigation.trends import holm
        raw = [0.001, 0.04, 0.03, 0.5]
        adjusted = holm(raw)
        self.assertTrue(all(a >= r for a, r in zip(adjusted, raw)))
        self.assertAlmostEqual(adjusted[0], 0.004)
        self.assertEqual(adjusted[3], 0.5)

    def test_change_labels_three_tiers(self):
        from tools.investigation.trends import label_changes
        rows = [{"change": {"points": -10.0, "p_value": 0.0001}}, {"change": {"points": 5.0, "p_value": 0.04}}, {"change": {"points": 1.0, "p_value": 0.7}}]
        label_changes(rows, "change")
        self.assertEqual(rows[0]["change"]["label"], "Clear improvement")
        self.assertTrue(rows[1]["change"]["label"].startswith("Possible worsening"))
        self.assertEqual(rows[2]["change"]["label"], "No clear change")

    def test_published_district_table_is_complete_and_consistent(self):
        from tools.investigation.trends import PUBLISHED_NEW_SERIES
        self.assertEqual(len(PUBLISHED_NEW_SERIES), 30)
        for name, (before, after, drop) in PUBLISHED_NEW_SERIES.items():
            self.assertAlmostEqual(before - after, drop, delta=0.15, msg=name)

    def test_eicv3_codes_must_map_one_to_one(self):
        from tools.investigation.trends import eicv3_names
        districts = {float(11 + index): f"D{index}" for index in range(30)}
        good = pd.DataFrame({"DISTRICT_eicv3": [f"{index:04d}" for index in range(30)] * 2, "district": [float(11 + index) for index in range(30)] * 2})
        self.assertEqual(len(eicv3_names(good, districts)), 30)
        bad = good.copy()
        bad.loc[0, "district"] = 12.0
        with self.assertRaises(ValueError):
            eicv3_names(bad, districts)

    def test_panel_mobility_counts_and_rejects_person_conflicts(self):
        from tools.investigation.trends import panel_mobility
        rows = []
        for household in range(60):
            poor3, poor4 = household % 2 == 0, household % 4 in (0, 1)
            for person in range(2):
                rows.append({"hhid": household, "caseid": household * 10 + person, "clust": household // 3, "province": 1.0, "pop_wt": 1.0, "POP_WT_eicv3": 1.0,
                             "pov_EICV3": 100 if poor3 else 0, "pov_EICV4": 1 if poor4 else 0})
        panel = pd.DataFrame(rows)
        result = panel_mobility(panel, {1.0: "Kigali City"})
        self.assertEqual(result["people"], 120)
        self.assertEqual(result["households"], 60)
        self.assertAlmostEqual(result["overall"]["escaped_poverty"]["pct"], 50.0)
        shares = sum(result["shares_of_people"][key]["pct"] for key in ("stayed_poor", "escaped", "fell_back", "never_poor"))
        self.assertAlmostEqual(shares, 100.0, delta=0.3)
        panel.loc[0, "pov_EICV4"] = 1 - panel.loc[1, "pov_EICV4"]
        with self.assertRaises(ValueError):
            panel_mobility(panel, {1.0: "Kigali City"})


class ProvenanceTests(unittest.TestCase):
    def test_every_result_is_labelled_and_unused_data_is_disclosed(self):
        from tools.investigation.provenance import data_used
        result = data_used({"source_file": "D:/x/EICV7 2023-2024.zip"}, {"financial_access": {"status": "Completed", "source_file": "D:/x/FinScope 2024.dta"},
                            "poverty_trends": {"status": "Completed", "files": {"EICV3": "a.sav", "EICV4": "b.sav", "EICV5": "c.dta", "panel": "p.dta"}}})
        self.assertTrue(all(item["linkage"].startswith(("Household-level linked", "Aggregate comparison only")) for item in result["used"]))
        self.assertTrue(any("FinScope 2012" in item["dataset"] for item in result["not_yet_used"]))
        finscope = next(item for item in result["used"] if item["dataset"].startswith("FinScope"))
        self.assertEqual(finscope["file"], "FinScope 2024.dta")
        self.assertTrue(finscope["linkage"].startswith("Aggregate comparison only"))


class FinScope2020Tests(unittest.TestCase):
    def frame(self, per_district=40):
        from tools.investigation import finance
        labels = ["Banked", "Other formal non-bank", "Informal ONLY", "Excluded"]
        rows = []
        for index, name in enumerate(sorted(finance.PUBLISHED_DISTRICTS)):
            for person in range(per_district):
                label = labels[(person + index) % 4]
                rows.append({"Individuals_Weight": 100.0 + person, "NUMVILL": f"{index}-{person // 4}", "A1": f"{['Eastern', 'Kigali City', 'Northern', 'Southern', 'Western'][index % 5]} Province", "A2": name, "A3": "Rural", "B1": 30.0,
                             "FAS": label, "Banked": 1.0 if label == "Banked" else np.nan, "Other_formal": 2.0 if label in ("Banked", "Other formal non-bank") and person % 2 == 0 or label == "Other formal non-bank" else np.nan,
                             "Informal": 3.0 if label in ("Informal ONLY", "Banked") and label != "Banked" or label == "Informal ONLY" else np.nan,
                             "FM4": "You are a registered mobile money user" if person % 3 else "You are not registered and have never used mobile money services",
                             "K1_2_1": "Yes" if label == "Banked" else "No"})
        return pd.DataFrame(rows)

    def frame2024(self, per_district=40):
        from tools.investigation import finance
        rows = [{"pop_wt": 100.0 + person, "a2": name, "f5": "You are a registered mobile money user" if person % 4 else "You have not heard about the Mobile Money"}
                for name in sorted(finance.PUBLISHED_DISTRICTS) for person in range(per_district)]
        return pd.DataFrame(rows)

    def test_valid_frame_passes_and_strands_partition_people(self):
        from tools.investigation.finscope2020 import STRANDS, prepare, validate
        frame = self.frame()
        validate(frame)
        prepared = prepare(frame)
        self.assertTrue((prepared[list(STRANDS)].sum(axis=1) == 1).all())
        self.assertTrue((prepared.formal == (prepared.banked | prepared.other_formal)).all())
        self.assertTrue(prepared.province.isin(["Eastern", "Kigali City", "Northern", "Southern", "Western"]).all())

    def test_inconsistent_strand_components_and_bad_inputs_are_rejected(self):
        from tools.investigation.finscope2020 import validate
        frame = self.frame()
        frame.loc[frame.FAS == "Excluded", "Informal"] = 3.0
        with self.assertRaises(ValueError):
            validate(frame)
        bad_weight = self.frame()
        bad_weight.loc[0, "Individuals_Weight"] = 0
        with self.assertRaises(ValueError):
            validate(bad_weight)
        unknown = self.frame()
        unknown.loc[0, "FAS"] = "Something else"
        with self.assertRaises(ValueError):
            validate(unknown)
        few = self.frame()
        few.loc[few.A2 == "Burera", "A2"] = "Gasabo"
        with self.assertRaises(ValueError):
            validate(few)

    def test_registered_mobile_money_excludes_unregistered_users(self):
        from tools.investigation import finance
        from tools.investigation.finscope2020 import REGISTERED_MOBILE_MONEY
        self.assertEqual(len(REGISTERED_MOBILE_MONEY), 3)
        self.assertNotIn("You use mobile money services, but you are not registered", REGISTERED_MOBILE_MONEY)
        self.assertIn("You use mobile money services, but you are not registered", finance.MOBILE_MONEY_USERS)

    def test_published_2024_values_carry_rounding_uncertainty(self):
        from tools.investigation.finscope2020 import published_2024
        values = published_2024("Nyarugenge", 400.0)
        self.assertEqual(values["excluded"]["published"], 0.0)
        self.assertEqual(values["formal"]["published"], 99)
        self.assertGreater(values["formal"]["se"], 0.0)
        self.assertGreater(values["excluded"]["se"], 0.0)

    def test_trend_rows_label_changes_with_neutral_words(self):
        from tools.investigation.finscope2020 import TREND_MEASURES, prepare, trend_rows
        names = sorted(__import__("tools.investigation.finance", fromlist=["x"]).PUBLISHED_DISTRICTS)
        rows = trend_rows(prepare(self.frame()), self.frame2024(), names)
        self.assertEqual(len(rows), 30)
        for row in rows:
            for measure in TREND_MEASURES:
                change = row[measure]["change"]
                self.assertIsNotNone(change)
                self.assertRegex(change["label"], r"^(Clear|Possible) (rise|fall)|^No clear change")
                self.assertNotIn("improvement", change["label"])

    def test_reconstruction_rules_follow_priority_and_ignore_mutuelle(self):
        from tools.investigation.finscope2020 import reconstruct, reconstruction_check, prepare
        base = self.frame()
        for column in ["K1_2_1", "KB2", "QF4_01", "QF5_01", "QF4_04", "QF5_04", "QF4_08", "QF4_05", "QF5_05", "QF4_02", "QF5_02", "QF4_06", "QF5_06", "QF4_07", "QF5_07", "GS9", "II3A", "I3A"]:
            base[column] = "No"
        base["FM4"] = "You have not heard about the Mobile Money"
        base["I3A"] = "Mutuelle de Sante (MdS)"
        self.assertTrue((reconstruct(base) == "excluded").all(), "Mutuelle alone must leave a person excluded")
        base.loc[0, "QF4_08"] = "Yes"
        base.loc[1, "QF4_05"] = "Yes"
        base.loc[2, "QF4_05"] = "Yes"
        base.loc[2, "KB2"] = "Yes"
        result = reconstruct(base)
        self.assertEqual([result[0], result[1], result[2]], ["informal_only", "other_formal", "banked"])
        check = reconstruction_check(prepare(base))
        self.assertEqual(set(check["rules"]), {"formal", "informal_only", "excluded"})
        self.assertEqual(set(check["difference_points"]), {"formal", "informal_only", "excluded"})

    def test_national_checks_flag_a_mismatch(self):
        from tools.investigation.finscope2020 import national_checks
        national = {"formal": {"pct": 77.0}, "informal_only": {"pct": 15.6}, "excluded": {"pct": 15.0}, "banked": {"pct": 36.2}}
        checks = {item["measure"]: item for item in national_checks(national, 62.3, 76.2)}
        self.assertTrue(checks["formal"]["ok"])
        self.assertFalse(checks["excluded"]["ok"])
        self.assertTrue(checks["registered_mobile_money_2024"]["ok"])


class FinScope2016Tests(unittest.TestCase):
    def frame(self, per_province=120):
        labels = ["Banked", "Other formal (non-bank)", "Informal only", "Excluded"]
        rows = []
        for index, province in enumerate(["Eastern", "Kigali City", "Northern", "Southern", "Western"]):
            for person in range(per_province):
                label = labels[(person + index) % 4]
                rows.append({"Individuals_Weight": 100.0 + person, "A1": f"{province} Province", "A2": float(person % 3 + 1), "AS_revised": label,
                             "Banked_popultaion": 1.0 if label == "Banked" else np.nan,
                             "other_formal_non_bank": 2.0 if label == "Other formal (non-bank)" else np.nan,
                             "Informal_revised": 3.0 if label == "Informal only" else np.nan})
        return pd.DataFrame(rows)

    def test_valid_frame_passes_and_strands_partition_people(self):
        from tools.investigation.finscope2016 import STRANDS, prepare, validate
        frame = self.frame()
        validate(frame)
        prepared = prepare(frame)
        self.assertTrue((prepared[list(STRANDS)].sum(axis=1) == 1).all())
        self.assertTrue(prepared.province.isin(["Eastern", "Kigali City", "Northern", "Southern", "Western"]).all())

    def test_inconsistent_or_bad_inputs_are_rejected(self):
        from tools.investigation.finscope2016 import validate
        excluded_with_use = self.frame()
        excluded_with_use.loc[excluded_with_use.AS_revised == "Excluded", "Informal_revised"] = 3.0
        with self.assertRaises(ValueError):
            validate(excluded_with_use)
        bad_weight = self.frame()
        bad_weight.loc[0, "Individuals_Weight"] = -1
        with self.assertRaises(ValueError):
            validate(bad_weight)
        missing_province = self.frame()
        missing_province = missing_province[~missing_province.A1.str.startswith("Western")]
        with self.assertRaises(ValueError):
            validate(missing_province)
        unknown = self.frame()
        unknown.loc[0, "AS_revised"] = "Other"
        with self.assertRaises(ValueError):
            validate(unknown)

    def test_province_changes_use_neutral_labels_and_the_right_family(self):
        from tools.investigation.finscope2016 import SERIES_MEASURES, PROVINCES, prepare, province_changes
        from tools.investigation.finscope2020 import prepare as prepare2020
        frame2020 = prepare2020(FinScope2020Tests("test_valid_frame_passes_and_strands_partition_people").frame(per_district=80))
        rows = province_changes(prepare(self.frame()), frame2020)
        self.assertEqual([row["province"] for row in rows], list(PROVINCES))
        for row in rows:
            for measure in SERIES_MEASURES:
                change = row[measure]["change"]
                self.assertIsNotNone(change)
                self.assertRegex(change["label"], r"^(Clear|Possible) (rise|fall)|^No clear change")
                self.assertNotIn("30 districts", change["label"])

    def test_national_checks_use_a_one_point_tolerance(self):
        from tools.investigation.finscope2016 import national_checks
        share = lambda value: {"p": value / 100, "se": 0.01, "n": 1000}
        national = {"banked": share(25.9), "other_formal": share(41.6), "informal_only": share(20.6), "excluded": share(14.0), "formal": share(67.5)}
        checks = {item["measure"]: item for item in national_checks(national)}
        self.assertTrue(checks["banked"]["ok"])
        self.assertFalse(checks["excluded"]["ok"])


class ProblemProfileTests(unittest.TestCase):
    def analysis(self, poverty=45.6, se=2.7, grid=39.1, grid_se=3.6, vup=18.2, vup_se=3.0, grade="Reliable"):
        def entry(name, value, standard_error, kind=grade):
            return {"indicator": name, "estimate_pct": value, "se_points": standard_error, "ci_lower": value - 2 * standard_error, "ci_upper": value + 2 * standard_error, "sample": 486, "grade": kind}
        return {"district": "Testville", "national": {"poverty": 27.4, "extreme_poverty": 5.4, "vup_among_poor": 15.9, "grid_connection": 50.8, "health_insurance": 85.3},
                "reliability": {"indicators": [entry("Poverty", poverty, se), entry("Extreme poverty", 5.5, 1.0), entry("VUP participation among poor households", vup, vup_se),
                                               entry("Grid electricity connection", grid, grid_se), entry("Current health insurance", 85.0, 1.8)]}}

    def test_z_is_signed_by_which_direction_is_worse(self):
        from tools.investigation.problem_profile import z_worse
        self.assertEqual(z_worse(45.6, 27.4, 2.7, "higher"), 6.74)
        self.assertEqual(z_worse(39.1, 50.8, 3.6, "lower"), 3.25)
        self.assertLess(z_worse(60.0, 50.0, 2.0, "lower"), 0)
        self.assertIsNone(z_worse(1.0, 2.0, 0.0, "higher"))
        self.assertIsNone(z_worse(1.0, 2.0, None, "higher"))

    def test_tiers_follow_the_stated_thresholds_and_low_reliability_cannot_stand_out(self):
        from tools.investigation.problem_profile import tier
        self.assertEqual(tier(2.0, "Reliable"), "stands_out")
        self.assertEqual(tier(1.5, "Reliable"), "possible")
        self.assertEqual(tier(0.9, "Reliable"), "in_line")
        self.assertEqual(tier(-2.0, "Reliable"), "better")
        self.assertEqual(tier(5.0, "Low reliability"), "possible")
        self.assertEqual(tier(None, "Reliable"), "context")

    def test_profile_ranks_by_z_and_reports_what_it_cannot_tell(self):
        from tools.investigation.problem_profile import build_profile
        profile = build_profile(self.analysis(), {"explained_excl_markers_share_pct": 6.0, "gap_points": 18.2, "unexplained_points": 12.0}, {"ranked_drivers": []}, None, None)
        self.assertEqual([row["id"] for row in profile["stands_out"]][:2], ["poverty", "grid_connection"])
        self.assertGreater(profile["stands_out"][0]["z"], profile["stands_out"][1]["z"])
        self.assertIn("main cause is not visible", " ".join(profile["flags"]))
        self.assertTrue(profile["cannot_tell"])
        self.assertIn("does not show causes", profile["method"])
        json.dumps(profile, allow_nan=False)

    def test_low_vup_gets_a_reporting_caveat_and_better_results_are_separate(self):
        from tools.investigation.problem_profile import build_profile
        low = build_profile(self.analysis(vup=8.0, vup_se=2.0), None, None, None, None)
        self.assertIn("reported participation", " ".join(low["flags"]))
        high = build_profile(self.analysis(vup=40.0, vup_se=3.0), None, None, None, None)
        self.assertIn("vup_among_poor", [row["id"] for row in high["better"]])
        self.assertNotIn("reported participation", " ".join(high["flags"]))

    def test_slow_progress_is_flagged_but_never_ranked(self):
        from tools.investigation.problem_profile import build_profile
        trends = {"status": "Completed", "selected": {"new_series": {}}, "new_series": [{"district": "Testville", "decrease_points": 3.6, "eicv5_modelled_pct": 46.9, "eicv7_published_pct": 43.3}]}
        profile = build_profile(self.analysis(), None, None, None, trends)
        self.assertEqual([row["id"] for row in profile["context"]], ["poverty_progress"])
        self.assertIsNone(profile["context"][0]["z"])
        self.assertIn("no district interval", " ".join(profile["flags"]))
        self.assertNotIn("poverty_progress", [row["id"] for row in profile["stands_out"]])

    def test_living_standards_numbers_reach_the_profile_and_its_flag(self):
        from tools.investigation.problem_profile import build_profile
        lens = {"shortfall_points": 12.0, "spending_difference_pct": -11.3, "explained_excl_markers_share_pct": 25.0, "explained_share_pct": 55.0,
                "explainability": {"living_standards_out_of_sample_r_squared": 0.58, "poverty_status_out_of_sample_r_squared": 0.25}}
        drivers = {"explained_excl_markers_share_pct": 11.0, "gap_points": 18.3, "unexplained_points": 12.0, "consumption_view": lens}
        profile = build_profile(self.analysis(), drivers, {"ranked_drivers": []}, None, None)
        living = profile["measured_factors"]["living_standards"]
        self.assertEqual(living["explained_share_pct"], 55.0)
        self.assertEqual(living["poverty_status_out_of_sample_r_squared"], 0.25)
        joined = " ".join(profile["flags"])
        self.assertIn("25% of the district's difference in typical spending", joined)
        self.assertIn("about 25% of the variation in who is poor", joined)
        json.dumps(profile, allow_nan=False)
        without = build_profile(self.analysis(), {"explained_excl_markers_share_pct": 11.0, "gap_points": 18.3, "unexplained_points": 12.0}, {"ranked_drivers": []}, None, None)
        self.assertIsNone(without["measured_factors"]["living_standards"])

    def test_headline_leads_with_the_unexplained_share_when_little_is_explained(self):
        from tools.investigation.diagnosis import diagnose
        drivers = {"district": "Testville", "gap_points": 18.3, "explained_excl_markers_share_pct": 6.0, "explained_share_pct": 34.0, "explained_points": 6.2, "unexplained_points": 12.0,
                   "groups": [{"id": "energy", "label": "Energy access and cooking fuel", "role": "access", "contribution_points": 1.3, "ci_lower": 0.5, "ci_upper": 2.1, "grade": "Strong association", "sign_stability": 99, "variables": []}],
                   "ranked_drivers": ["energy"], "protective_factors": [], "marker_groups": [], "profile": [], "place_effect": None, "not_measured": []}
        result = diagnose({"selected": {"district": "Testville", "poverty": 40, "sample_households": 486}, "national": {"poverty": 27.4}}, drivers, {"sources": []})
        self.assertTrue(result["headline"].startswith("Only 6% of Testville's 18.3-point difference from Rwanda"), result["headline"])
        self.assertIn("Strongest measured factor: energy access", result["headline"])


class UsedColumnsTests(unittest.TestCase):
    @staticmethod
    def read_calls(module):
        """Every (table, columns) pair a module passes to _read(archive, TABLE, [columns])."""
        import ast
        import importlib
        source = Path(module.__file__).read_text(encoding="utf-8")
        found = []
        for node in ast.walk(ast.parse(source)):
            if isinstance(node, ast.Call) and getattr(node.func, "id", "") == "_read" and len(node.args) >= 3 and isinstance(node.args[2], ast.List):
                table_arg = node.args[1]
                path = getattr(module, table_arg.id, None) if isinstance(table_arg, ast.Name) else None
                if isinstance(path, str):
                    found.append((Path(path).stem.lower(), [item.value for item in node.args[2].elts if isinstance(item, ast.Constant)]))
        return found

    def test_every_column_the_eicv7_modules_read_is_listed(self):
        from tools.investigation import analysis, drivers, used_columns
        from tools.investigation.catalog import SENSITIVE
        calls = self.read_calls(analysis) + self.read_calls(drivers)
        self.assertGreaterEqual(len(calls), 8)
        for table, columns in calls:
            for column in columns:
                if SENSITIVE.search(column):
                    continue
                self.assertIn(column.lower(), used_columns.USED_COLUMNS.get(table, set()), f"{table}.{column} is read but not listed as used")

    def test_finscope_modules_read_only_listed_columns(self):
        import re
        from tools.investigation import finance, finscope2016, finscope2020, used_columns
        self.assertLessEqual({name.lower() for name in finance.CORE_COLUMNS + finance.OPTIONAL_COLUMNS}, used_columns.FINSCOPE[used_columns.FINSCOPE_2024])
        self.assertEqual(finance.OPTIONAL_PATTERNS, used_columns.FINSCOPE_2024_PATTERNS)
        self.assertLessEqual({name.lower() for name in finscope2020.COLUMNS}, used_columns.FINSCOPE[used_columns.FINSCOPE_2020])
        import inspect
        read = set(re.findall(r'"([A-Z][A-Z0-9_]*)"', inspect.getsource(finscope2020.reconstruct)))
        self.assertTrue(read, "reconstruct should read named columns")
        self.assertLessEqual({name.lower() for name in read}, used_columns.FINSCOPE_2020_CHECK, "reconstruct reads a column that is not listed")
        self.assertLessEqual({name.lower() for name in finscope2016.COLUMNS}, used_columns.FINSCOPE[used_columns.FINSCOPE_2016])
        self.assertIn("qf4_05", used_columns.FINSCOPE[used_columns.FINSCOPE_2020])

    def test_trend_readers_use_only_listed_columns(self):
        import ast
        from tools.investigation import trends, used_columns
        listed = set().union(*[used_columns.USED_COLUMNS[key] for key in used_columns.TREND_FILES])
        tree = ast.parse(Path(trends.__file__).read_text(encoding="utf-8"))
        seen = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Call):
                for keyword in node.keywords:
                    if keyword.arg in ("usecols", "columns") and isinstance(keyword.value, ast.List):
                        seen.update(item.value for item in keyword.value.elts if isinstance(item, ast.Constant))
        self.assertTrue(seen)
        for column in seen:
            if column.lower() in {"hhid"}:
                continue
            self.assertIn(column.lower(), listed, f"{column} is read by the trend module but not listed as used")

    def test_catalog_marks_listed_columns_used_and_others_not(self):
        from tools.investigation.catalog import classify
        self.assertEqual(classify("s10cq4", "Savings", "CS_S10C_Savings", "EICV7", 0, 100)[0], "Used")
        self.assertNotEqual(classify("s10cq1", "Savings", "CS_S10C_Savings", "EICV7", 0, 100)[0], "Used")
        self.assertEqual(classify("QF4_05", "Registered in own name", "Rwanda FinScope 2020_Final Data", "FinScope", 0, 100)[0], "Used")
        self.assertEqual(classify("qf3_05", "Account", "T2 Finance and Poverty - FinScope 2024 Financial Inclusion", "FinScope", 0, 100)[0], "Used")
        self.assertNotEqual(classify("e18_1", "Borrowed", "T2 Finance and Poverty - FinScope 2024 Financial Inclusion", "FinScope", 0, 100)[0], "Used")


class EducationWorkTests(unittest.TestCase):
    COLUMNS = ["hhid", "s1q3y", "s4aq1", "s4aq2", "s4aq7", "s4bq4", "s4bq6", "s6aq2", "s6aq4", "s6aq5", "s6aq6", "s6aq7", "s6bq7", "s6bq8"]

    def person(self):
        own = "Own account worker (without regular employees)"
        rows = [
            (1, 35, "Yes", "Primary", np.nan, "Yes", "Yes", "No", "No", "Yes", "No", "No", "Employee", "Daily worker"),
            (1, 33, "No", np.nan, np.nan, "No", np.nan, "No", "No", "No", "No", "No", np.nan, np.nan),
            (1, 10, "Yes", "Primary", "Yes", "Yes", "Yes", "No", "No", "No", "No", "No", np.nan, np.nan),
            (2, 40, "Yes", "Lower secondary", np.nan, "Yes", "Yes", "Yes", "No", "No", "No", "No", own, np.nan),
            (2, 8, "Yes", "Primary", "No", np.nan, np.nan, "No", "No", "No", "No", "No", np.nan, np.nan),
            (3, 70, "Yes", "Primary", np.nan, "Yes", "Yes", "No", "No", "No", "No", "Yes", "Contributing family worker", np.nan),
            (4, 24, "Yes", "Upper secondary", np.nan, "Yes", "Yes", "No", "No", "No", "No", "No", np.nan, np.nan),
            (4, 28, "Yes", "Primary", np.nan, "No", np.nan, "No", "No", "No", "No", "No", np.nan, np.nan),
        ]
        frame = pd.DataFrame(rows, columns=self.COLUMNS)
        frame["weight"] = 1.0
        return frame

    def test_household_variables_follow_the_definitions(self):
        from tools.investigation.drivers import education_work_variables
        result = education_work_variables(self.person(), pd.Index([1, 2, 3, 4], name="hhid"))
        self.assertEqual(result.illiterate_adult_share.round(2).tolist(), [0.5, 0.0, 0.0, 0.5])
        self.assertEqual(result.no_secondary_adult.tolist(), [1.0, 0.0, 1.0, 0.0])
        self.assertEqual(result.child_out_of_school.tolist(), [0.0, 1.0, 0.0, 0.0])
        self.assertEqual(result.casual_work.tolist(), [1.0, 0.0, 0.0, 0.0])
        self.assertEqual(result.unpaid_family_work.tolist(), [0.0, 0.0, 1.0, 0.0])
        self.assertEqual(result.no_adult_worker.tolist(), [0.0, 0.0, 0.0, 1.0], "a household with no working-age adult is not counted as one whose adults lack a job")
        self.assertFalse(result.isna().any().any())

    def test_people_who_cannot_read_are_not_literate_even_though_they_are_not_asked_to_write(self):
        from tools.investigation.drivers import education_work_variables
        person = self.person()
        self.assertTrue(person.loc[1, "s4bq6"] != person.loc[1, "s4bq6"], "fixture: the write question is blank for a non-reader")
        result = education_work_variables(person, pd.Index([1, 2, 3, 4], name="hhid"))
        self.assertEqual(result.illiterate_adult_share.iloc[0], 0.5)

    def test_validation_reports_values_and_rejects_definitions_that_miss_nisr(self):
        from unittest.mock import patch as patch_object
        from tools.investigation import drivers
        loose = {name: (0.0, 1000.0) for name in drivers.NISR_PERSON_FIGURES}
        with patch_object.dict(drivers.NISR_PERSON_FIGURES, loose):
            checks = {item["measure"]: item for item in drivers.validate_education_work(self.person())}
        self.assertEqual(checks["literacy_15_plus"]["calculated"], 66.7)
        self.assertEqual(checks["literacy_15_24"]["calculated"], 100.0)
        self.assertEqual(checks["no_formal_education_10_plus"]["calculated"], 14.3)
        self.assertEqual(checks["secondary_10_plus"]["calculated"], 28.6)
        self.assertEqual(checks["primary_10_plus"]["calculated"], 57.1)
        self.assertEqual(checks["university_10_plus"]["calculated"], 0.0)
        with self.assertRaises(ValueError):
            drivers.validate_education_work(self.person())

    def test_validation_rejects_questions_that_are_not_answered_where_expected(self):
        from unittest.mock import patch as patch_object
        from tools.investigation import drivers
        loose = {name: (0.0, 1000.0) for name in drivers.NISR_PERSON_FIGURES}
        missing_reading = self.person()
        missing_reading.loc[2, "s4bq4"] = np.nan
        no_main_job = self.person()
        no_main_job.loc[0, "s6bq7"] = np.nan
        for broken in (missing_reading, no_main_job):
            with patch_object.dict(drivers.NISR_PERSON_FIGURES, loose), self.assertRaises(ValueError):
                drivers.validate_education_work(broken)

    def test_the_six_new_variables_are_in_the_model_with_the_right_groups(self):
        from tools.investigation.drivers import VARIABLES
        groups = {name: group for name, _, group, _, _ in VARIABLES}
        for name in ("illiterate_adult_share", "no_secondary_adult", "child_out_of_school"):
            self.assertEqual(groups[name], "education")
        for name in ("no_adult_worker", "casual_work", "unpaid_family_work"):
            self.assertEqual(groups[name], "livelihood")
        self.assertEqual(len(VARIABLES), 40)


class NarrativeTests(unittest.TestCase):
    findings = {"findings": [{"id": "poverty"}]}
    context = {"findings": [{"id": "poverty", "value": 45.6, "national": 27.4}]}

    def review(self, text, ids=("poverty",)):
        from tools.investigation.ai import validate_review
        return validate_review({"finding_ids": ["poverty"], "option_titles": [], "question_ids": [], "narrative": [{"text": text, "finding_ids": list(ids)}]},
                               self.findings, {"options": []}, self.context)

    def test_narrative_with_known_numbers_is_kept(self):
        result = self.review("Poverty is 45.6% against 27.4% nationally, an association only.")
        self.assertEqual(result["narrative"][0]["finding_ids"], ["poverty"])

    def test_invented_number_is_dropped_but_valid_sentences_stay(self):
        from tools.investigation.ai import validate_review
        value = {"finding_ids": ["poverty"], "option_titles": [], "question_ids": [], "narrative": [
            {"text": "Poverty is 51.2% in this district.", "finding_ids": ["poverty"]}, {"text": "Poverty is 45.6% against 27.4% nationally.", "finding_ids": ["poverty"]}]}
        result = validate_review(value, self.findings, {"options": []}, self.context)
        self.assertEqual([item["text"] for item in result["narrative"]], ["Poverty is 45.6% against 27.4% nationally."])
        self.assertIn("51.2", result["dropped_statements"][0])

    def test_causal_language_is_dropped(self):
        result = self.review("Low grid access caused poverty of 45.6%.")
        self.assertNotIn("narrative", result)
        self.assertIn("causation", result["dropped_statements"][0])

    def test_unknown_finding_reference_is_dropped(self):
        result = self.review("Poverty is 45.6%.", ids=("invented",))
        self.assertNotIn("narrative", result)
        self.assertEqual(result["status"], "Completed")

    def test_narrative_without_evidence_context_is_rejected(self):
        from tools.investigation.ai import validate_review
        with self.assertRaises(ValueError):
            validate_review({"finding_ids": ["poverty"], "option_titles": [], "question_ids": [], "narrative": [{"text": "Poverty is 45.6%.", "finding_ids": ["poverty"]}]},
                            self.findings, {"options": []})



class PriorityBoardTests(unittest.TestCase):
    @staticmethod
    def profile(clear, possible=()):
        row = lambda name, z: {"id": name, "label": name.title(), "district_pct": 40.0, "rwanda_pct": 30.0, "difference_points": 10.0, "z": z, "reliability": "Reliable", "source": "EICV7"}
        return {"profile": {"headline": "h", "stands_out": [row(*item) for item in clear], "possible": [row(*item) for item in possible]}, "run_id": "r", "run_at": "t"}

    def test_districts_rank_by_total_size_of_clear_differences(self):
        from tools.investigation.board import build_board
        profiles = {"A": self.profile([("poverty", 2.5)]), "B": self.profile([("poverty", 3.0), ("power", 2.1)]), "C": self.profile([], [("poverty", 1.5)]), "D": self.profile([("poverty", 4.0)]),
                    "F": self.profile([("poverty", 20.0)])}
        board = build_board(profiles, ["A", "B", "C", "D", "E", "F"])
        self.assertEqual([row["district"] for row in board["districts"]], ["F", "B", "D", "A", "C"])
        self.assertEqual(board["districts"][0]["score"], 6.0)
        self.assertEqual(board["missing"], ["E"])
        self.assertEqual(board["signals"][0]["id"], "poverty")
        self.assertEqual(board["signals"][0]["districts"][0]["district"], "F")
        self.assertIn("not a ranking of need", board["note"])

    def test_board_is_built_from_saved_runs_newest_completed_run_wins(self):
        from tools.investigation.board import write_board
        def save(store, run_id, district, status, stamp, clear):
            folder = store / "runs" / run_id
            folder.mkdir(parents=True)
            (folder / "manifest.json").write_text(json.dumps({"id": run_id, "district": district, "status": status, "updated_at": stamp}), encoding="utf-8")
            (folder / "08-preparing-findings.json").write_text(json.dumps({"result": {"problem_profile": self.profile(clear)["profile"]}}), encoding="utf-8")
        with tempfile.TemporaryDirectory() as directory:
            store = Path(directory)
            save(store, "old", "Kamonyi", "Completed", "2026-01-01", [("poverty", 2.1)])
            save(store, "new", "Kamonyi", "Warning", "2026-02-01", [("poverty", 5.0)])
            save(store, "bad", "Kamonyi", "Failed", "2026-03-01", [("poverty", 9.0)])
            save(store, "other", "Huye", "Completed", "2026-01-05", [])
            board = write_board(store, ["Kamonyi", "Huye", "Gasabo"])
            self.assertEqual([row["district"] for row in board["districts"]], ["Kamonyi", "Huye"])
            self.assertEqual(board["districts"][0]["run_id"], "new")
            self.assertEqual(board["missing"], ["Gasabo"])
            self.assertEqual(json.loads((store / "priority_board.json").read_text(encoding="utf-8"))["districts_included"], 2)

    def test_batch_runs_do_not_send_figures_to_the_ai_provider(self):
        from tools.investigation.ai import providers_from_env
        with patch.dict("os.environ", {"GROQ_API_KEY": "gsk_x", "INVESTIGATION_NO_AI": "1"}):
            self.assertEqual(providers_from_env(), [])


class FailLoudlyTests(unittest.TestCase):
    def test_low_scratch_space_stops_the_run_with_a_plain_message(self):
        with patch("tools.investigation.catalog.shutil.disk_usage") as usage:
            usage.return_value.free = 300 * 1024 ** 2
            with self.assertRaisesRegex(OSError, "Free at least 2 GB"):
                catalog.check_scratch_space()
            usage.return_value.free = 5 * 1024 ** 3
            catalog.check_scratch_space()

    def test_disk_full_failures_are_recognised_and_other_failures_are_not(self):
        tables = [{"filename": "a", "status": "Failed", "error": "OSError: [Errno 28] No space left on device"},
                  {"filename": "b", "status": "Failed", "error": "ValueError: bad file"}, {"filename": "c", "status": "Completed"}]
        self.assertEqual([item["filename"] for item in catalog.disk_full_tables(tables)], ["a"])

    def test_run_stops_and_saves_no_catalog_when_the_disk_fills_during_discovery(self):
        import uuid
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "Data").mkdir()
            (root / "Data" / "tiny.csv").write_text("a,b" + chr(10) + "1,2", encoding="utf-8")
            full = {"tables": [{"filename": "big.zip!/x.dta", "status": "Failed", "error": "OSError: [Errno 28] No space left on device"}]}
            with patch.object(engine, "STORE", root / "store"), patch.object(engine, "DATA", root / "Data"), patch.object(engine, "check_scratch_space"),                     patch.object(engine, "discover", return_value=full), patch.object(engine, "refresh_usage", side_effect=lambda value: value):
                work = engine.Investigation(str(uuid.uuid4()), "Gisagara")
                with self.assertRaisesRegex(OSError, "Nothing was saved"):
                    work.run()
                self.assertEqual(work.manifest["status"], "Failed")
                self.assertEqual(list((root / "store" / "cache").glob("catalog-*.json")), [])


class FinScope2012Tests(unittest.TestCase):
    def frame(self):
        names = finscope2012.COLUMNS
        rows = [
            {"banked": True, "other": False, "group": False},
            {"banked": False, "other": True, "group": False},
            {"banked": False, "other": False, "group": True},
            {"banked": False, "other": False, "group": False},
        ] * 25
        frame = pd.DataFrame({name: "No" for name in names}, index=range(len(rows)))
        frame[finscope2012.WEIGHT] = [100.0 + index for index in range(len(rows))]
        for name in finscope2012.BANK_ANSWERED + finscope2012.OTHER_FORMAL_ANSWERED:
            frame[name] = np.nan
        for index, item in enumerate(rows):
            if item["banked"]:
                frame.loc[index, "J3A"] = "Yes"
            if item["other"]:
                frame.loc[index, "J8A"] = "Yes"
            if item["group"]:
                frame.loc[index, "K1A"] = "Yes"
        return frame

    def test_strands_partition_people_and_follow_the_framework_order(self):
        prepared = finscope2012.prepare(self.frame())
        self.assertTrue((prepared[list(finscope2012.STRANDS)].sum(axis=1) == 1).all())
        self.assertEqual([int(prepared[key].sum()) for key in finscope2012.STRANDS], [25, 25, 25, 25])
        both = self.frame()
        both.loc[0, ["J3A", "K1A", "J8A"]] = "Yes"
        self.assertTrue(finscope2012.prepare(both).loc[0, "banked"])

    def test_bad_inputs_are_rejected(self):
        bad_weight = self.frame()
        bad_weight.loc[0, finscope2012.WEIGHT] = -1
        with self.assertRaises(ValueError):
            finscope2012.validate(bad_weight)
        with self.assertRaises(ValueError):
            finscope2012.validate(self.frame().drop(columns=["J3A"]))

    def test_checks_use_a_one_point_tolerance(self):
        share = lambda value: {"p": value / 100, "se": 0.01, "n": 1000}
        rebuilt = {"banked": share(23.8), "other_formal": share(19.0), "informal_only": share(30.0), "excluded": share(39.0), "formal": share(42.0)}
        found = {item["measure"]: item for item in finscope2012.checks(rebuilt)}
        self.assertTrue(found["banked"]["ok"])
        self.assertFalse(found["excluded"]["ok"])
        self.assertEqual(found["excluded"]["difference_points"], 11.0)

    def test_real_file_gap_is_reported_and_series_stays_without_2012(self):
        root = Path(__file__).resolve().parents[2] / "Data"
        if not list(root.rglob(finscope2012.FILE_PATTERN)):
            self.skipTest("FinScope 2012 file is not on this machine")
        result = finscope2012.summary(root)
        self.assertFalse(result["agrees_with_published"])
        self.assertEqual({item["measure"] for item in result["checks"]}, set(finscope2012.PUBLISHED_2012))
        from tools.investigation import finscope2016
        self.assertNotIn("2012", [item["key"] for item in finscope2016.financial_series(root)["series"]])



class VariableChecksTests(unittest.TestCase):
    @staticmethod
    def recorded():
        from tools.investigation import variable_checks
        return [{"table": table, "columns": list(columns), "measure": measure, "published": published, "estimate": estimate, "tolerance": tolerance,
                 "ok": abs(estimate - published) <= tolerance, "source": source}
                for table, columns, measure, published, estimate, tolerance, source in variable_checks.RECORDED_RESULTS]

    def test_recorded_results_cover_the_checks_and_include_failures(self):
        from tools.investigation import variable_checks
        results = self.recorded()
        self.assertEqual([item["measure"] for item in results], [check["measure"] for check in variable_checks.CHECKS])
        self.assertEqual([item["published"] for item in results], [check["published"] for check in variable_checks.CHECKS])
        self.assertGreaterEqual(len(results), 60)
        self.assertTrue(any(not item["ok"] for item in results), "failed checks must stay on record, not be dropped")

    def test_only_columns_whose_checks_all_pass_are_marked_used(self):
        from tools.investigation import used_columns, variable_checks
        from tools.investigation.analysis import HOUSEHOLD, PERSON
        passing = variable_checks.validated_columns(self.recorded())
        tables = {"household": Path(HOUSEHOLD).stem.lower(), "person": Path(PERSON).stem.lower()}
        for kind, table in tables.items():
            listed = used_columns.EICV7_VALIDATED[table]
            self.assertLessEqual(listed, passing[kind], f"{table}: a column is marked Used without passing every check")
            for column in passing[kind]:
                self.assertIn(column, used_columns.USED_COLUMNS[table], f"{table}.{column} passed its checks but is not listed")
        for kind, column in variable_checks.NOT_VALIDATED:
            self.assertNotIn(column, used_columns.USED_COLUMNS[tables[kind]], f"{column} failed its check and must stay unvalidated")
        failed = {(item["table"], column) for item in self.recorded() if not item["ok"] for column in item["columns"]}
        self.assertLessEqual({(kind, column) for kind, column in variable_checks.NOT_VALIDATED}, failed)
        self.assertGreaterEqual(sum(len(columns) for columns in used_columns.EICV7_VALIDATED.values()), 30)

    def test_every_column_the_checks_read_is_used_or_recorded_as_not_validated(self):
        from tools.investigation import used_columns, variable_checks
        from tools.investigation.analysis import HOUSEHOLD, PERSON
        for kind, path, columns in (("household", HOUSEHOLD, variable_checks.HOUSEHOLD_COLUMNS), ("person", PERSON, variable_checks.PERSON_COLUMNS)):
            for column in columns:
                self.assertTrue(column in used_columns.USED_COLUMNS[Path(path).stem.lower()] or (kind, column) in variable_checks.NOT_VALIDATED, f"{kind}.{column}")
        for check in variable_checks.CHECKS:
            read = set(variable_checks.HOUSEHOLD_COLUMNS if check["table"] == "household" else variable_checks.PERSON_COLUMNS)
            self.assertLessEqual(set(check["columns"]), read, check["measure"])

    def test_catalog_marks_validated_columns_used_and_failed_ones_not(self):
        from tools.investigation.catalog import classify
        self.assertEqual(classify("s5aq1", "Type of habitat", "CS_S01_S5_S7_Household", "EICV7", 0, 100)[0], "Used")
        self.assertEqual(classify("s4bq14i", "Writing a computer program", "CS_S0_S1_S2_S3_S4_S6A_S6B_S6C_Person", "EICV7", 0, 100)[0], "Used")
        self.assertNotEqual(classify("s6bq4", "ISIC level 1", "CS_S0_S1_S2_S3_S4_S6A_S6B_S6C_Person", "EICV7", 0, 100)[0], "Used")
        self.assertNotEqual(classify("s1q12", "Father still alive", "CS_S0_S1_S2_S3_S4_S6A_S6B_S6C_Person", "EICV7", 0, 100)[0], "Used")

    def test_checks_compare_a_weighted_estimate_with_the_published_figure(self):
        from tools.investigation import variable_checks
        household = pd.DataFrame({"weight": [1.0, 3.0, 1.0, 1.0], "s5aq1": ["Umudugudu (Planned rural settlement)", "Dispersed/Isolated housing", "Umudugudu (Planned rural settlement)", "Modern planned urban area"]})
        check = next(item for item in variable_checks.CHECKS if item["published"] == 67.9 and item["columns"] == ("s5aq1",))
        self.assertAlmostEqual(variable_checks._share(household, check["mask"](household)), 100 * 2 / 6)
        bad = dict(household, weight=[1.0, -1.0, 1.0, 1.0])
        with self.assertRaises(ValueError):
            variable_checks._share(pd.DataFrame(bad), check["mask"](pd.DataFrame(bad)))

    @unittest.skipUnless(list(Path("Data").rglob("*EICV7 2023-2024.zip")), "EICV7 archive not available")
    def test_recorded_results_match_a_fresh_run(self):
        from tools.investigation import variable_checks
        fresh = variable_checks.calculate(Path("Data"))
        for item, old in zip(fresh, self.recorded()):
            self.assertAlmostEqual(item["estimate"], old["estimate"], places=2, msg=item["measure"])
            self.assertEqual(item["ok"], old["ok"], item["measure"])

class FileLockTests(unittest.TestCase):
    def test_a_briefly_locked_manifest_is_saved_after_a_retry(self):
        real = os.replace
        calls = []

        def locked_twice(source, target):
            calls.append(1)
            if len(calls) <= 2:
                raise PermissionError(5, "Access is denied")
            return real(source, target)

        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "manifest.json"
            with patch("tools.investigation.fileio.os.replace", side_effect=locked_twice), patch("tools.investigation.fileio.time.sleep"):
                engine.write_json(target, {"status": "Running"})
            self.assertEqual(json.loads(target.read_text(encoding="utf-8")), {"status": "Running"})
            self.assertEqual(len(calls), 3)
            self.assertEqual([item.name for item in Path(directory).iterdir()], ["manifest.json"])

    def test_a_permanently_locked_file_still_fails_loudly_and_leaves_no_temporary_file(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "manifest.json"
            with patch("tools.investigation.fileio.os.replace", side_effect=PermissionError(5, "Access is denied")), patch("tools.investigation.fileio.time.sleep"):
                with self.assertRaises(PermissionError):
                    engine.write_json(target, {"status": "Running"})
            self.assertEqual(list(Path(directory).iterdir()), [])


if __name__ == "__main__":
    unittest.main()

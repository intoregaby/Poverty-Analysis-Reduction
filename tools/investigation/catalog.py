"""Inventory every local data file and variable without exporting respondent records."""

from __future__ import annotations

import hashlib
import copy
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from typing import Callable

import pandas as pd

from .used_columns import USED_COLUMNS, USED_PATTERNS

SUPPORTED = {".csv", ".xlsx", ".sav", ".dta", ".json", ".parquet"}
TOPICS = {
    "poverty": ("pov", "consumption", "expenditure", "spending"),
    "employment": ("employ", "work", "job", "occupation", "wage"),
    "income": ("income", "business", "enterprise", "profit"),
    "agriculture": ("agric", "crop", "livestock", "land", "farm"),
    "finance": ("bank", "saving", "credit", "loan", "mobile money", "finscope"),
    "housing": ("water", "sanitation", "toilet", "electric", "housing", "roof"),
    "education": ("school", "literacy", "education", "training"),
    "health": ("health", "insurance", "illness"),
    "social_protection": ("vup", "support", "benefit", "transfer", "ubudehe"),
}
SENSITIVE = re.compile(r"(^|_)(hhid|personid|pid|nid|national_id|phone|account|name|birth_date|dob|key)(_|$)", re.I)
GEO = re.compile(r"district|province|sector|cell|village|urban|rural|strata|cluster|^ea$|^psu$", re.I)
WEIGHT = re.compile(r"(^|_)(weight|pop_wt|hh_wt|pers_wt|wt)(_|$)|sampling weight", re.I)


def version(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


MINIMUM_FREE_BYTES = 2 * 1024 ** 3
DISK_FULL = re.compile(r"No space left|\[Errno 28\]|ENOSPC", re.I)


def check_scratch_space(minimum: int = MINIMUM_FREE_BYTES) -> None:
    """Archives are unpacked into the temporary folder; stop early instead of cataloguing half the data."""
    folder = tempfile.gettempdir()
    free = shutil.disk_usage(folder).free
    if free < minimum:
        raise OSError(f"Only {free / 1024 ** 3:.1f} GB is free where temporary files go ({folder}). Free at least {minimum / 1024 ** 3:.0f} GB or point TEMP and TMP at a drive with more space, then run again.")


def disk_full_tables(tables: list[dict]) -> list[dict]:
    return [table for table in tables if table.get("status") == "Failed" and DISK_FULL.search(table.get("error") or "")]


def find_seven_zip() -> str | None:
    """Locate 7-Zip: SEVEN_ZIP variable, PATH, Program Files, then D:/application."""
    candidates = [os.environ.get("SEVEN_ZIP"), shutil.which("7z"), shutil.which("7zz")]
    for root in (os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)")):
        if root:
            candidates.append(str(Path(root) / "7-Zip" / "7z.exe"))
    if Path("D:/application").exists():
        candidates += [str(path) for path in Path("D:/application").glob("*Zip*/7-Zip/7z.exe")]
    return next((item for item in candidates if item and Path(item).is_file()), None)


def extract_member(archive_path: Path, archive: zipfile.ZipFile, member: zipfile.ZipInfo, target_path: Path) -> None:
    """Copy one archive member to target_path. If Python's zipfile cannot read it
    (for example Deflate64), retry through 7-Zip and check the size."""
    try:
        with archive.open(member) as source, target_path.open("wb") as target:
            shutil.copyfileobj(source, target, 1024 * 1024)
        return
    except (NotImplementedError, RuntimeError, zipfile.BadZipFile):
        pass
    seven_zip = find_seven_zip()
    if not seven_zip:
        raise RuntimeError("Archive member needs 7-Zip (not found; set the SEVEN_ZIP variable)")
    with target_path.open("wb") as target:
        operation = subprocess.run([seven_zip, "e", "-so", str(archive_path), member.filename],
                                   stdout=target, stderr=subprocess.PIPE, check=False)
    if operation.returncode:
        raise RuntimeError(f"7-Zip could not extract member (exit {operation.returncode})")
    if target_path.stat().st_size != member.file_size:
        raise RuntimeError("7-Zip output size differs from the archive listing (possible truncation)")


def survey_identity(name: str) -> tuple[str | None, str | None]:
    if re.search(r"EICV\s*3", name, re.I) and re.search(r"EICV\s*4", name, re.I) and "panel" in name.lower():
        return "EICV3–EICV4 panel", "2010/11 and 2013/14; wave-specific period requires validation"
    match = re.search(r"EICV\s*([34567])", name, re.I)
    if match:
        periods = {"3": "2010/11", "4": "2013/14", "5": "2016/17", "6": "2019/20", "7": "2023/24"}
        return f"EICV{match.group(1)}", periods[match.group(1)]
    match = re.search(r"FinScope\s*(20\d{2})", name, re.I)
    if match:
        return "FinScope", match.group(1)
    return None, None


def is_used(table: str, name: str) -> bool:
    key = table.lower()
    return name.lower() in USED_COLUMNS.get(key, set()) or any(re.match(pattern, name, re.I) for pattern in USED_PATTERNS.get(key, []))


def classify(name: str, label: str, table: str, survey: str | None, missing: int | None, rows: int | None) -> tuple[str, str, list[str]]:
    combined = f"{name} {label}".lower()
    topics = [topic for topic, words in TOPICS.items() if any(word in combined for word in words)]
    if SENSITIVE.search(name) or re.search(r"(household|respondent|person)\s+(id|identifier|number)|(?:number|identifier)\s+(?:of\s+the\s+(?:selected\s+)?)?(household|respondent|person)|(?:account|phone|identity)\s+number", label, re.I):
        return "Restricted", "Potential record identifier or personal field; aggregate-only use requires an approved join.", topics
    if missing is not None and rows and missing == rows:
        return "Relevant but unavailable" if topics else "Unreliable", "Every observed value is missing.", topics
    if missing is not None and rows and missing / rows > 0.8:
        return "Unreliable", "More than 80% of observed values are missing; do not analyze without review.", topics
    if survey in {"EICV3", "EICV4", "EICV5", "EICV3–EICV4 panel"} and "pov" in combined:
        return "Not comparable", "Older poverty measures cannot be directly compared with EICV7; use NISR's modelled comparison.", topics
    if is_used(table, name):
        return "Used", "This exact table and variable are read by a validated aggregate analysis module.", topics
    if GEO.search(name) or GEO.search(label) or WEIGHT.search(name) or WEIGHT.search(label):
        return "Requires processing", "Geographic or survey-design field; validate codes and survey identity before use.", topics
    if topics:
        return "Requires processing", "Topically relevant; definition, denominator, weight, and linkage need validation.", topics
    return "Not relevant", "No validated link to the current district questions was established.", topics


def refresh_usage(catalog: dict) -> dict:
    """Reclassify cached metadata without rereading unchanged survey records."""
    for table in catalog["tables"]:
        survey, period = survey_identity(table.get("container") or table["filename"])
        table.update({"survey": survey, "survey_period": period, "collection_period": period,
                      "survey_identity_basis": "Inferred from the file/container name; wave and collection dates require source-document validation"})
        if survey == "EICV3–EICV4 panel":
            table["relationships"] = []
        for variable in table.get("variables", []):
            classification, reason, topics = classify(variable["name"], variable.get("label", ""), table.get("table") or "",
                                                       table.get("survey"), variable.get("missing_values"), table.get("row_count"))
            variable.update({"classification": classification, "classification_reason": reason, "topics": topics})
    catalog["classification_version"] = version(Path(__file__))
    return catalog


def _tabular_metadata(path: Path, suffix: str):
    if suffix in {".sav", ".dta"}:
        import pyreadstat

        reader = pyreadstat.read_sav if suffix == ".sav" else pyreadstat.read_dta
        def scan(encoding):
            _, meta = reader(str(path), metadataonly=True, encoding=encoding)
            names = list(meta.column_names)
            labels = dict(zip(names, meta.column_labels or []))
            types = meta.readstat_variable_types or {}
            missing = Counter({name: 0 for name in names})
            rows = 0
            if meta.number_rows is None or meta.number_rows < 0:
                frames = [reader(str(path), encoding=encoding)[0]]
            else:
                frames = (reader(str(path), row_offset=offset, row_limit=20000, encoding=encoding)[0]
                          for offset in range(0, meta.number_rows, 20000))
            for frame in frames:
                rows += len(frame)
                missing.update(frame.isna().sum().to_dict())
            return [(name, labels.get(name) or "", str(types.get(name, "unknown")), int(missing[name])) for name in names], rows, []
        try:
            return scan(None)
        except UnicodeDecodeError:
            return scan("latin1")

    if suffix == ".csv":
        reader = pd.read_csv(path, chunksize=20000, low_memory=False)
        missing = Counter()
        types = {}
        rows = 0
        names = []
        for frame in reader:
            names = list(frame.columns)
            rows += len(frame)
            missing.update(frame.isna().sum().to_dict())
            for name, dtype in frame.dtypes.items():
                types.setdefault(name, str(dtype))
        if not names:
            names = list(pd.read_csv(path, nrows=0).columns)
        return [(name, "", types.get(name, "unknown"), int(missing[name])) for name in names], rows, []

    if suffix == ".parquet":
        import pyarrow.parquet as pq

        parquet = pq.ParquetFile(path)
        names = parquet.schema_arrow.names
        types = {field.name: str(field.type) for field in parquet.schema_arrow}
        missing = Counter({name: 0 for name in names})
        rows = 0
        for batch in parquet.iter_batches(batch_size=20000):
            frame = batch.to_pandas()
            rows += len(frame)
            missing.update(frame.isna().sum().to_dict())
        return [(name, "", types[name], int(missing[name])) for name in names], rows, []

    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        records = payload if isinstance(payload, list) else [payload]
        frame = pd.json_normalize(records)
        return [(name, "", str(frame[name].dtype), int(frame[name].isna().sum())) for name in frame.columns], len(frame), []

    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        tables = []
        for sheet in workbook:
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator, ())
            names = [str(value).strip() if value is not None else f"unnamed_{index + 1}" for index, value in enumerate(header)]
            missing = Counter({name: 0 for name in names})
            types = {}
            rows = 0
            for cells in iterator:
                rows += 1
                for name, value in zip(names, cells):
                    if value is None or value == "":
                        missing[name] += 1
                    else:
                        types.setdefault(name, type(value).__name__)
            tables.append((sheet.title, [(name, "", types.get(name, "unknown"), int(missing[name])) for name in names], rows))
        workbook.close()
        return [], 0, tables

    raise ValueError(f"Unsupported format: {suffix}")


def _record(path: Path, display: str, parent: str, version_id: str, survey: str | None, period: str | None, member: str | None = None) -> list[dict]:
    suffix = path.suffix.lower()
    basic = {
        "id": hashlib.sha256(display.encode()).hexdigest()[:16],
        "filename": display,
        "container": parent,
        "member": member,
        "format": suffix.lstrip(".") or "unknown",
        "source": "User-supplied local file; publisher is not independently proven by filename",
        "survey": survey,
        "survey_period": period,
        "collection_period": period if survey else None,
        "geographic_coverage": "Unknown until geographic variables are validated",
        "version": version_id,
        "access": "Restricted microdata; aggregate outputs only" if suffix in {".sav", ".dta"} else "Local file; review license before redistribution",
        "relationships": [],
    }
    if suffix not in SUPPORTED:
        return [{**basic, "table": Path(display).stem, "status": "Skipped", "error": "Unsupported file format or archive metadata.", "row_count": None, "variables": []}]
    try:
        variables, rows, sheets = _tabular_metadata(path, suffix)
        items = sheets if sheets else [(Path(display).stem, variables, rows)]
        result = []
        for table_name, columns, row_count in items:
            annotated = []
            for name, label, dtype, missing in columns:
                classification, explanation, topics = classify(name, label, table_name, survey, missing, row_count)
                annotated.append({
                    "name": name, "label": label, "data_type": dtype, "missing_values": missing,
                    "weight": bool(WEIGHT.search(name) or WEIGHT.search(label)),
                    "geographic_identifier": bool(GEO.search(name) or GEO.search(label)),
                    "classification": classification, "classification_reason": explanation,
                    "topics": topics,
                })
            result.append({**basic, "id": hashlib.sha256(f"{display}:{table_name}".encode()).hexdigest()[:16], "table": table_name,
                           "status": "Completed", "error": None, "row_count": row_count, "variables": annotated})
        return result
    except Exception as error:
        return [{**basic, "table": Path(display).stem, "status": "Failed", "error": f"{type(error).__name__}: {error}", "row_count": None, "variables": []}]


def discover(root: Path, progress: Callable[[dict], None] | None = None, previous: dict | None = None) -> dict:
    tables = []
    reusable = {}
    for item in (previous or {}).get("tables", []):
        reusable.setdefault(item.get("container"), []).append(item)
    files = sorted(path for path in root.rglob("*") if path.is_file())
    for index, path in enumerate(files, start=1):
        relative = path.relative_to(root).as_posix()
        digest = version(path)
        survey, period = survey_identity(relative)
        cached = reusable.get(relative, [])
        if cached and all(item.get("version", "").split(":")[0] == digest for item in cached):
            tables.extend(copy.deepcopy(cached))
            if progress:
                progress({"file": relative, "table": f"{path.name} (unchanged content; catalog reused)", "files_discovered": index,
                          "tables_inventoried": len(tables), "variables_inventoried": sum(len(t["variables"]) for t in tables)})
            continue
        if path.suffix.lower() == ".zip":
            try:
                with zipfile.ZipFile(path) as archive:
                    for member in archive.infolist():
                        if member.is_dir():
                            continue
                        display = f"{relative}!/{member.filename}"
                        member_version = f"{digest}:{member.CRC:08x}:{member.file_size}"
                        if member.filename.startswith("__MACOSX/") or Path(member.filename).name.startswith("._"):
                            tables.extend(_record(path, display, relative, member_version, survey, period, member.filename))
                        elif Path(member.filename).suffix.lower() in SUPPORTED:
                            with tempfile.TemporaryDirectory(prefix="nisr-catalog-") as temporary:
                                extracted = Path(temporary) / f"table{Path(member.filename).suffix.lower()}"
                                try:
                                    extract_member(path, archive, member, extracted)
                                    tables.extend(_record(extracted, display, relative, member_version, survey, period, member.filename))
                                except Exception as error:
                                    tables.append({"filename": display, "container": relative, "member": member.filename,
                                                   "format": Path(member.filename).suffix.lstrip("."), "survey": survey,
                                                   "survey_period": period, "version": member_version, "table": Path(member.filename).stem,
                                                   "status": "Failed", "error": f"{type(error).__name__}: {error}", "variables": []})
                        else:
                            tables.extend(_record(path, display, relative, member_version, survey, period, member.filename))
                        if progress:
                            progress({"file": relative, "table": member.filename, "files_discovered": index,
                                      "tables_inventoried": len(tables), "variables_inventoried": sum(len(t["variables"]) for t in tables)})
            except Exception as error:
                tables.append({"filename": relative, "table": None, "format": "zip", "status": "Failed", "error": str(error), "variables": []})
        else:
            tables.extend(_record(path, relative, relative, digest, survey, period))
            if progress:
                progress({"file": relative, "table": path.name, "files_discovered": index,
                          "tables_inventoried": len(tables), "variables_inventoried": sum(len(t["variables"]) for t in tables)})

    by_survey = {}
    for table in tables:
        if table.get("survey"):
            by_survey.setdefault((table["survey"], table.get("survey_period")), []).append(table)
    for group in by_survey.values():
        names_by_table = {table.get("id"): {v["name"].lower() for v in table.get("variables", [])} for table in group}
        for table in group:
            current = names_by_table[table.get("id")]
            shared = [other.get("id") for other in group if other is not table and (current & names_by_table[other.get("id")] & {"hhid", "personid", "pid"})]
            table["relationships"] = shared

    return {
        "data_root": str(root.resolve()),
        "files_discovered": len(files),
        "tables_inventoried": len(tables),
        "variables_inventoried": sum(len(table.get("variables", [])) for table in tables),
        "failed_tables": sum(table.get("status") == "Failed" for table in tables),
        "skipped_tables": sum(table.get("status") == "Skipped" for table in tables),
        "tables": tables,
    }
